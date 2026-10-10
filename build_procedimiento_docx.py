import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_shading(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="B0C4DE"):
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
        <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
        <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
        <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
        <w:left w:val="none"/>
        <w:right w:val="none"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def build_procedimiento_docx():
    doc = docx.Document()

    # Colors
    c_primary = RGBColor(7, 55, 99)     # 073763
    c_secondary = RGBColor(28, 69, 135) # 1C4587
    c_accent = RGBColor(46, 117, 181)   # 2E75B5
    c_dark = RGBColor(30, 30, 30)

    # Margins
    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)

    # HEADER TABLE (Membrete Institucional)
    tbl_hdr = doc.add_table(rows=2, cols=3)
    tbl_hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_hdr, "073763")

    # Column widths
    col_widths = [Inches(1.8), Inches(3.6), Inches(1.8)]
    for row in tbl_hdr.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    # Logo in Cell (0,0)
    c0 = tbl_hdr.cell(0, 0)
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists("Logo Optimus Health.jpeg"):
        try:
            p0.add_run().add_picture("Logo Optimus Health.jpeg", width=Inches(1.4))
        except Exception:
            p0.add_run("OPTIMUS HEALTH IPS").bold = True

    # Title in Cell (0,1)
    c1 = tbl_hdr.cell(0, 1)
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run("SISTEMA DE GESTIÓN DE LA CALIDAD\nPROGRAMA NEWCARE • CRÓNICOS COMPLEJOS\nPROCEDIMIENTO DE SELECCIÓN Y BAREMACIÓN")
    r1.bold = True
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = c_primary

    # Metadata in Cell (0,2)
    c2 = tbl_hdr.cell(0, 2)
    p2 = c2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r2 = p2.add_run("CÓDIGO: TH-PRO-004\nVERSIÓN: 2.0\nVIGENCIA: 2026 - 2028\nPÁGINA: 1 de 18")
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = c_secondary

    # Row 1 spans
    c_sub = tbl_hdr.cell(1, 0)
    c_sub.merge(tbl_hdr.cell(1, 1)).merge(tbl_hdr.cell(1, 2))
    p_sub = c_sub.paragraphs[0]
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("OPTIMUS HEALTH IPS • SEDE PRINCIPAL SINCELEJO, SUCRE • METODOLOGÍA FUNDACIÓN NEW HEALTH")
    r_sub.font.size = Pt(8)
    r_sub.bold = True
    r_sub.font.color.rgb = RGBColor(255, 255, 255)
    set_cell_shading(c_sub, "073763")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # TITULO PRINCIPAL
    h1 = doc.add_paragraph()
    h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_h1 = h1.add_run("PROCEDIMIENTO OPERATIVO ESTÁNDAR (POE)\nSELECCIÓN, BAREMACIÓN Y DIMENSIONAMIENTO DEL TALENTO HUMANO ASISTENCIAL")
    r_h1.bold = True
    r_h1.font.size = Pt(16)
    r_h1.font.color.rgb = c_primary

    sub1 = doc.add_paragraph()
    sub1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub1 = sub1.add_run("Modelo de Atención Integral en Insuficiencias Orgánicas Crónicas Complejas (IOCC)\nAlineado con el Método NEWPALEX® de la Fundación New Health y el Sistema Único de Habilitación")
    r_sub1.font.size = Pt(11)
    r_sub1.italic = True
    r_sub1.font.color.rgb = c_secondary
    sub1.paragraph_format.space_after = Pt(18)

    # =========================================================================
    # 1. OBJETIVO Y ALCANCE
    # =========================================================================
    h_sec1 = doc.add_paragraph()
    r_sec1 = h_sec1.add_run("1. PROPÓSITO Y ALCANCE")
    r_sec1.bold = True
    r_sec1.font.size = Pt(13)
    r_sec1.font.color.rgb = c_primary

    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_obj.add_run("El presente procedimiento tiene como ").font.color.rgb = c_dark
    p_obj.add_run("propósito ").bold = True
    p_obj.add_run("establecer directrices claras, reproducibles y objetivas para la atracción, selección, baremación de idoneidad, contratación e inducción del talento humano asistencial que conforma el equipo tratante del Programa NewCare en Optimus Health IPS (Sincelejo, Sucre). Este proceso garantiza que cada profesional y técnico cuente con las competencias científico-clínicas, la calidez humana, la resiliencia psicológica y el perfil ético requeridos para la atención de pacientes con Insuficiencias Orgánicas Crónicas Complejas (Falla Cardíaca, EPOC, Insuficiencia Renal Crónica y Cirrosis Hepática) y necesidades de cuidados paliativos.")

    p_alc = doc.add_paragraph()
    p_alc.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_alc.add_run("El ").font.color.rgb = c_dark
    p_alc.add_run("alcance ").bold = True
    p_alc.add_run("abarca a todo el personal asistencial intramural y extramural (domiciliario) vinculado o por vincular al programa: Director Médico / Líder NewCare, Médicos Especialistas consultores, Médicos Generales de seguimiento, Gestor de Casos (Case Manager), Profesionales de Enfermería, Auxiliares de Enfermería, Psicología, Trabajo Social, Nutrición y Profesionales de Rehabilitación (Fisioterapia y Fonoaudiología).")

    # =========================================================================
    # 2. MARCO NORMATIVO Y REFERENCIAL
    # =========================================================================
    h_sec2 = doc.add_paragraph()
    r_sec2 = h_sec2.add_run("2. MARCO NORMATIVO Y CONCEPTUAL")
    r_sec2.bold = True
    r_sec2.font.size = Pt(13)
    r_sec2.font.color.rgb = c_primary

    normas = [
        ("Resolución 3100 de 2019 (Ministerio de Salud y Protección Social):", "Define los estándares y criterios del Sistema Único de Habilitación de prestadores de servicios de salud en Colombia, en particular el Estándar de Talento Humano para atención ambulatoria, extramural y hospitalaria."),
        ("Ley 1733 de 2014 (Ley Consuelo Devis Saavedra):", "Reglamenta los servicios de cuidados paliativos para el manejo integral de pacientes con enfermedades terminales, crónicas, degenerativas e irreversibles, exigiendo idoneidad profesional y trabajo interdisciplinario."),
        ("Ley 1164 de 2007 (Talento Humano en Salud):", "Establece disposiciones relacionadas con los procesos de formación, ejercicio y desempeño del talento humano en salud, así como la obligatoriedad del ReTHUS."),
        ("Estándares de Acreditación NewCare (Fundación New Health):", "Requisito 2.4.2 y Capítulo 3: Exige la definición formal y baremada de los criterios de selección, valoración curricular (50%), pruebas psicométricas (20%) y entrevista por competencias (30%).")
    ]

    for tit, desc in normas:
        p_n = doc.add_paragraph(style='List Bullet')
        p_n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r_t = p_n.add_run(tit + " ")
        r_t.bold = True
        r_t.font.color.rgb = c_secondary
        r_d = p_n.add_run(desc)
        r_d.font.color.rgb = c_dark

    # =========================================================================
    # 3. METODOLOGÍA DE SELECCIÓN Y BAREMACIÓN (100 PUNTOS)
    # =========================================================================
    h_sec3 = doc.add_paragraph()
    r_sec3 = h_sec3.add_run("3. METODOLOGÍA OPERATIVA DE SELECCIÓN Y BAREMACIÓN")
    r_sec3.bold = True
    r_sec3.font.size = Pt(13)
    r_sec3.font.color.rgb = c_primary

    p_fases = doc.add_paragraph()
    p_fases.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_fases.add_run("El proceso de selección en Optimus Health IPS se divide en tres fases secuenciales y obligatorias:")

    fases_list = [
        ("Fase 1: Filtros Excluyentes de Habilitación (SI/NO):", "Verificación legal y documental obligatoria de requisitos mínimos. Si un candidato presenta un solo incumplimiento, queda automáticamente descartado del proceso."),
        ("Fase 2: Baremación Triádica Objetiva (100 Puntos):", "Evaluación ponderada distribuida en: Valoración de Hoja de Vida / CV (50%), Evaluación Psicotécnica (20%) y Entrevista Semiestructurada por Competencias (30%)."),
        ("Fase 3: Dictamen de Idoneidad, Contratación e Inducción:", "Emisión de calificación final, asignación de rol en el organigrama y formulación del Plan Individual de Inducción y Entrenamiento NewCare.")
    ]
    for f_tit, f_desc in fases_list:
        p_f = doc.add_paragraph(style='List Bullet')
        r_ft = p_f.add_run(f_tit + " ")
        r_ft.bold = True
        r_ft.font.color.rgb = c_accent
        p_f.add_run(f_desc).font.color.rgb = c_dark

    # TABLA DE FILTROS EXCLUYENTES
    h_tbl_flt = doc.add_paragraph()
    h_tbl_flt.add_run("Tabla 1. Criterios de Selección Habilitantes Mínimos (Filtros SI / NO)").bold = True
    h_tbl_flt.runs[0].font.color.rgb = c_primary

    tbl_flt = doc.add_table(rows=10, cols=4)
    tbl_flt.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_flt, "073763")

    flt_headers = ["Código", "Requisito Excluyente Habilitante", "Soporte Documental Exigido", "Criterio"]
    for i, h in enumerate(flt_headers):
        cell = tbl_flt.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(cell, "073763")

    flt_rows = [
        ("FLT-01", "Título Profesional Universitario oficial expedido por IES en Colombia o convalidado.", "Diploma y Acta de Grado oficial", "Obligatorio (SÍ/NO)"),
        ("FLT-02", "Inscripción activa en Registro Único Nacional del Talento Humano en Salud (ReTHUS).", "Certificado de verificación SISPRO", "Obligatorio (SÍ/NO)"),
        ("FLT-03", "Tarjeta Profesional expedida por el colegio profesional respectivo vigente.", "Carné o resolución de colegiatura", "Obligatorio (SÍ/NO)"),
        ("FLT-04", "Certificado de Formación en Violencia Sexual (Res. 459/2012) mínimo 40 horas.", "Certificación académica válida", "Obligatorio (SÍ/NO)"),
        ("FLT-05", "Soporte Vital Básico (BLS) o Avanzado (ACLS) vigente (menor a 2 años).", "Certificado avalado (AHA/IES)", "Obligatorio (SÍ/NO)"),
        ("FLT-06", "Ausencia de sanciones ético-disciplinarias, penales, fiscales o de procuraduría.", "Certificados de antecedentes vigentes", "Obligatorio (SÍ/NO)"),
        ("FLT-07", "Experiencia asistencial clínica mínima certificada de 1 año en adultos/crónicos.", "Certificaciones laborales formales", "Obligatorio (SÍ/NO)"),
        ("FLT-08", "Disponibilidad para desplazamientos domiciliarios en Sincelejo y área de cobertura.", "Cláusula operativa en contrato", "Obligatorio (SÍ/NO)"),
        ("FLT-09", "Compromiso firmado de participación en la Certificación NewCare / New Health.", "Carta de compromiso suscrita", "Obligatorio (SÍ/NO)")
    ]

    for r_idx, r_data in enumerate(flt_rows, start=1):
        for c_idx, val in enumerate(r_data):
            c = tbl_flt.cell(r_idx, c_idx)
            c.text = val
            p = c.paragraphs[0]
            p.runs[0].font.size = Pt(9)
            if c_idx == 0 or c_idx == 3:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.runs[0].bold = True
            if r_idx % 2 == 0:
                set_cell_shading(c, "F7F9FC")

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # TABLA DE BAREMACIÓN DE CV (50 PUNTOS)
    h_tbl_cv = doc.add_paragraph()
    h_tbl_cv.add_run("Tabla 2. Baremo Oficial de Hoja de Vida / Currículum Vitae (50 Puntos Máximos)").bold = True
    h_tbl_cv.runs[0].font.color.rgb = c_primary

    tbl_cv = doc.add_table(rows=10, cols=5)
    tbl_cv.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_cv, "073763")

    cv_headers = ["Dimensión Evaluada", "Peso NewCare", "Puntos Máx", "Criterio de Asignación de Puntaje", "Soporte Exigido"]
    for i, h in enumerate(cv_headers):
        cell = tbl_cv.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(cell, "073763")

    cv_rows_doc = [
        ("1. Experiencia Laboral", "30%", "15.0 pts", "General: 1.0 pt/año (máx 7.0 pts) | Específica en Crónicos/Paliativos/Domicilio: 2.0 pts/año (máx 8.0 pts)", "Certificaciones laborales de IPS"),
        ("2. Formación Postgrado", "12%", "6.0 pts", "Especialidad médica/Maestría clínica: 4.0 pts | Diplomado formal en Paliativos/Dolor (>100h): 2.0 pts", "Diploma y acta oficial IES"),
        ("3. Centros de Referencia", "14%", "7.0 pts", "Pasantía o rotación clínica en unidad de Cuidados Paliativos o Red New Health: 3.5 pts/mes (máx 7.0 pts)", "Certificado formal del centro"),
        ("4. Actividad Docente", "11%", "5.5 pts", "Docencia universitaria en salud: 1.5 pts/semestre (máx 3.0 pts) | Docente de cuidadores/comunidad: 1.0 pt/curso (máx 2.5 pts)", "Certificación universitaria"),
        ("5. Publicaciones", "9%", "4.5 pts", "Artículo en revista indexada: 2.5 pts/artículo (máx 4.5 pts) | Capítulo de libro o guía clínica: 1.5 pts/publicación", "Copia artículo DOI / ISBN"),
        ("6. Investigación", "10%", "5.0 pts", "Investigador en grupo avalado MinCiencias o semillero institucional: 2.5 pts/proyecto activo (máx 5.0 pts)", "Certificado CvLAC / MinCiencias"),
        ("7. Idiomas", "4%", "2.0 pts", "Dominio certificado de inglés técnico (Nivel B2 o superior): 2.0 pts", "Certificado examen oficial"),
        ("8. Gestión y Calidad", "10%", "5.0 pts", "Experiencia demostrada en coordinación de servicios, auditoría médica o gestión de casos: 2.5 pts/año (máx 5.0 pts)", "Certificado de cargo de dirección"),
        ("TOTAL BAREMO CV", "100%", "50.0 pts", "Sumatoria ponderada de expediente curricular documental", "Expediente auditado completo")
    ]

    for r_idx, r_data in enumerate(cv_rows_doc, start=1):
        for c_idx, val in enumerate(r_data):
            c = tbl_cv.cell(r_idx, c_idx)
            c.text = val
            p = c.paragraphs[0]
            p.runs[0].font.size = Pt(8.5)
            is_tot = "TOTAL" in r_data[0]
            if is_tot:
                p.runs[0].bold = True
                set_cell_shading(c, "E8EEF5")
            elif c_idx in [1, 2]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.runs[0].bold = True
            if r_idx % 2 == 0 and not is_tot:
                set_cell_shading(c, "F7F9FC")

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # TABLA DE PSICOTÉCNICO Y ENTREVISTA (50 PUNTOS)
    h_tbl_psi = doc.add_paragraph()
    h_tbl_psi.add_run("Tabla 3. Baremo Psicotécnico (20 Pts) y Entrevista por Competencias (30 Pts)").bold = True
    h_tbl_psi.runs[0].font.color.rgb = c_primary

    tbl_psi = doc.add_table(rows=13, cols=4)
    tbl_psi.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_psi, "073763")

    psi_headers = ["Componente / Variable", "Instrumento / Factor Evaluado", "Criterio de Aprobación", "Puntaje"]
    for i, h in enumerate(psi_headers):
        cell = tbl_psi.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(cell, "073763")

    psi_rows_doc = [
        ("Test Psicotécnico CPR (60%)", "Personalidad Resistente: Compromiso, Control interno y Reto adaptativo", "Percentil >= 65 en subescalas", "12.0 pts"),
        ("Test Psicotécnico DAS (40%)", "Escala de Actitudes Disfuncionales: Aprobación, autonomía y límites terapéuticos", "Percentil < 40 en disfuncionalidad", "8.0 pts"),
        ("Subtotal Psicotécnico", "Evaluación psicométrica de estabilidad emocional y resiliencia al Burnout", "Mínimo 14.0 pts para idoneidad", "20.0 pts"),
        ("Entrevista Dim. 1: Motivación", "Vocación por el alivio del sufrimiento y la dignidad en cronicidad avanzada", "Escala de 0 a 10 (Nivel Alto 8-10)", "3.33 pts"),
        ("Entrevista Dim. 2: Trabajo en Equipo", "Sinergia interdisciplinaria, respeto interprofesional y liderazgo horizontal", "Escala de 0 a 10 (Nivel Alto 8-10)", "3.33 pts"),
        ("Entrevista Dim. 3: Manejo de Crisis", "Templanza y mediación serena ante crisis familiares y descompensaciones", "Escala de 0 a 10 (Nivel Alto 8-10)", "3.33 pts"),
        ("Entrevista Dim. 4: Trato Humano", "Calidez, compasión clínica y validación incondicional del paciente", "Escala de 0 a 10 (Nivel Alto 8-10)", "3.33 pts"),
        ("Entrevista Dim. 5: Escucha Activa", "Decodificación empática en las 4 dimensiones (física, emocional, social, espiritual)", "Escala de 0 a 10 (Nivel Alto 8-10)", "3.33 pts"),
        ("Entrevista Dim. 6: Malas Noticias", "Dominio del protocolo SPIKES/EPICEE y verdad soportable", "Escala de 0 a 10 (Nivel Alto 8-10)", "3.34 pts"),
        ("Entrevista Dim. 7: Juicio Clínico", "Seguridad técnica, autonomía responsable y toma de decisiones bioéticas", "Escala de 0 a 10 (Nivel Alto 8-10)", "3.34 pts"),
        ("Entrevista Dim. 8: Calidad y Mejora", "Apertura a la auditoría continua, aprendizaje y mejora en la práctica clínica", "Escala de 0 a 10 (Nivel Alto 8-10)", "3.33 pts"),
        ("Entrevista Dim. 9: Cuidador y Comunidad", "Sensibilidad a la sobrecarga del cuidador (Zarit) y activación de redes solidarias", "Escala de 0 a 10 (Nivel Alto 8-10)", "3.34 pts"),
    ]

    for r_idx, r_data in enumerate(psi_rows_doc, start=1):
        for c_idx, val in enumerate(r_data):
            c = tbl_psi.cell(r_idx, c_idx)
            c.text = val
            p = c.paragraphs[0]
            p.runs[0].font.size = Pt(8.5)
            is_sub = "Subtotal" in r_data[0]
            if is_sub:
                p.runs[0].bold = True
                set_cell_shading(c, "E8EEF5")
            elif c_idx == 3:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.runs[0].bold = True
            if r_idx % 2 == 0 and not is_sub:
                set_cell_shading(c, "F7F9FC")

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # TABLA DE DICTAMEN FINAL
    h_dic = doc.add_paragraph()
    h_dic.add_run("Tabla 4. Escala de Dictamen Final de Idoneidad y Ponderación").bold = True
    h_dic.runs[0].font.color.rgb = c_primary

    tbl_dic = doc.add_table(rows=4, cols=3)
    tbl_dic.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_dic, "073763")

    dic_headers = ["Puntaje Consolidado", "Dictamen Institucional", "Acción de Calidad / Asignación en el Programa"]
    for i, h in enumerate(dic_headers):
        cell = tbl_dic.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(cell, "073763")

    dic_rows = [
        (">= 85.0 Puntos", "IDÓNEO SOBRESALIENTE", "Asignación directa e inmediata al rol asistencial y liderazgo de caso NewCare."),
        ("70.0 - 84.9 Puntos", "IDÓNEO CONDICIONADO", "Apto para contratación con Plan de Formación Continuada obligatorio en los primeros 90 días."),
        ("< 70.0 Puntos", "NO IDÓNEO", "No alcanza el perfil mínimo exigido por los estándares de la Fundación New Health. Candidato no seleccionado.")
    ]

    for r_idx, r_data in enumerate(dic_rows, start=1):
        for c_idx, val in enumerate(r_data):
            c = tbl_dic.cell(r_idx, c_idx)
            c.text = val
            p = c.paragraphs[0]
            p.runs[0].font.size = Pt(9)
            if c_idx in [0, 1]:
                p.runs[0].bold = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if r_idx == 1:
                set_cell_shading(c, "E2F0D9")
            elif r_idx == 2:
                set_cell_shading(c, "FFF2CC")
            else:
                set_cell_shading(c, "FCE4D6")

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # =========================================================================
    # 4. ROLES, COMPETENCIAS Y EQUIPO ASIGNADO EN OPTIMUS HEALTH
    # =========================================================================
    h_sec4 = doc.add_paragraph()
    r_sec4 = h_sec4.add_run("4. PERFILES, COMPETENCIAS Y TALENTO HUMANO DESIGNADO EN OPTIMUS HEALTH")
    r_sec4.bold = True
    r_sec4.font.size = Pt(13)
    r_sec4.font.color.rgb = c_primary

    p_perfiles = doc.add_paragraph()
    p_perfiles.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_perfiles.add_run("A continuación se consolida la estructura formal del talento humano de Optimus Health IPS asignado a la operativización del Programa NewCare, de conformidad con el Organigrama Asistencial revisado y las exigencias de la Fundación New Health:")

    # TABLA DE ROLES Y NOMBRES
    tbl_roles = doc.add_table(rows=14, cols=4)
    tbl_roles.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_roles, "073763")

    roles_headers = ["Cargo / Rol en el Programa", "Profesional Designado", "Nivel de Complejidad", "Funciones Principales en el Modelo NewCare"]
    for i, h in enumerate(roles_headers):
        cell = tbl_roles.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(cell, "073763")

    roles_data = [
        ("Director Médico / Líder NewCare", "Dr. Julio Herrera", "Transversal (Media y Alta)", "Liderazgo técnico, direccionamiento estratégico, articulación con aseguradoras y supervisión de auditorías."),
        ("Coordinadora Asistencial", "Enf. Dayana Solipaz", "Transversal Operativa", "Coordinación operativa de agendas, gestión de recursos, monitoreo de indicadores y cumplimiento de visitas."),
        ("Gestor de Casos (Case Manager)", "Enf. Maria Isabel Borja", "Transversal (Media y Alta)", "Monitoreo de cohorte activa, seguimiento proactivo, transiciones hospital-domicilio y enlace continuo familiar."),
        ("Médico Especialista Internista", "Dr. Jhair Martinez", "Alta Complejidad", "Manejo de multimorbilidad avanzada, ajuste terapéutico complejo, interconsultas y Junta de Cronicidad."),
        ("Médico Especialista Intensivista", "Dr. Pedro Herrera", "Alta Complejidad", "Manejo de descompensaciones agudas graves, limitación del esfuerzo terapéutico y bioética clínica."),
        ("Médico Especialista Pediatra", "Dra. Janna Arboleda", "Alta Complejidad Pediátrica", "Atención médica especializada y acompañamiento de pacientes crónicos complejos pediátricos."),
        ("Médico General Asistencial", "Médicos Tratantes de IPS", "Básica y Media Complejidad", "Valoración multidimensional, formulación y actualización del PAI, titulación y manejo domiciliario."),
        ("Auxiliares de Enfermería", "Edgardo Navarro, Luis Carlos Ricardo, Luis Sampayo, Nubia Solorzano, Maria Alejandra Martinez", "Básica y Media Complejidad", "Cuidado directo en domicilio, signos vitales, administración de fármacos, educación al cuidador y alerta temprana."),
        ("Fisioterapeuta", "Mara Salcedo", "Básica y Media Complejidad", "Rehabilitación física y respiratoria según PAI, prevención de atrofia y pautas de movilización al cuidador."),
        ("Fonoaudiólogo", "Melisa Aldana", "Media y Alta Complejidad", "Rehabilitación en deglución y comunicación, prevención de broncoaspiración y pautas de alimentación segura."),
        ("Psicólogo", "Angélica Monterroza", "Transversal", "Acompañamiento emocional, manejo del impacto diagnóstico, claudicación del cuidador (Zarit) y duelo."),
        ("Nutricionista", "Sara Chavez", "Media y Alta Complejidad", "Valoración nutricional (MNA), soporte dietario en insuficiencias orgánicas y educación nutricional familiar."),
        ("Trabajador Social", "Tatiana Guerra", "Transversal", "Diagnóstico sociofamiliar (Gijón), activación de redes de apoyo, gestión comunitaria y asamblea de pacientes.")
    ]

    for r_idx, r_data in enumerate(roles_data, start=1):
        for c_idx, val in enumerate(r_data):
            c = tbl_roles.cell(r_idx, c_idx)
            c.text = val
            p = c.paragraphs[0]
            p.runs[0].font.size = Pt(8.5)
            if c_idx in [0, 1]:
                p.runs[0].bold = True
            if c_idx == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if r_idx % 2 == 0:
                set_cell_shading(c, "F7F9FC")

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # =========================================================================
    # 5. FIRMAS DE APROBACIÓN
    # =========================================================================
    h_firmas = doc.add_paragraph()
    h_firmas.add_run("5. CONTROL DE EMISIÓN, REVISIÓN Y APROBACIÓN").bold = True
    h_firmas.runs[0].font.color.rgb = c_primary

    tbl_firmas = doc.add_table(rows=4, cols=3)
    tbl_firmas.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_firmas, "073763")

    firmas_hdr = ["ELABORÓ", "REVISÓ", "APROBÓ"]
    for i, h in enumerate(firmas_hdr):
        cell = tbl_firmas.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(cell, "073763")

    firmas_data = [
        ("Nombre: Enf. Dayana Solipaz\nCargo: Coordinadora Asistencial / Calidad\nFecha: 08/10/2026",
         "Nombre: Dr. Julio Herrera\nCargo: Director Médico / Líder NewCare\nFecha: 08/10/2026",
         "Nombre: Gerencia General\nCargo: Representante Legal Optimus Health IPS\nFecha: 08/10/2026"),
        ("Firma: _________________________", "Firma: _________________________", "Firma: _________________________"),
        ("Estado: Aprobado para Implementación", "Estado: Aprobado para Auditoría New Health", "Estado: Vigente para Operación")
    ]

    for r_idx, r_data in enumerate(firmas_data, start=1):
        for c_idx, val in enumerate(r_data):
            c = tbl_firmas.cell(r_idx, c_idx)
            c.text = val
            p = c.paragraphs[0]
            p.runs[0].font.size = Pt(8.5)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if r_idx == 3:
                p.runs[0].bold = True
                set_cell_shading(c, "E8EEF5")

    output_doc = "PROCEDIMIENTO_Y_BAREMO_SELECCION_TALENTO_HUMANO_NEWCARE_OPTIMUS_HEALTH.docx"
    doc.save(output_doc)
    print(f"Documento formal guardado exitosamente en: {output_doc}")

if __name__ == "__main__":
    build_procedimiento_docx()
