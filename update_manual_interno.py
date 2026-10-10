import docx
import shutil
import os

def update_manual():
    src_file = "Manual NEWCARE INTERNO - Optimus Health .docx"
    backup_file = "Manual NEWCARE INTERNO - Optimus Health (Copia Seguridad Pre-Baremacion).docx"

    if not os.path.exists(backup_file):
        shutil.copy2(src_file, backup_file)
        print(f"Copia de seguridad creada en: {backup_file}")
    else:
        print("Copia de seguridad ya existía.")

    doc = docx.Document(src_file)

    # Let's inspect paragraphs and update placeholders
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt == "En los anexos se encuentra el documento":
            p.text = (
                "El detalle metodológico, las tablas de ponderación, los criterios psicotécnicos (CPR y DAS), "
                "la rúbrica de entrevista y el simulador de cálculo automático se encuentran formalizados en el documento institucional "
                "'TH-PRO-004: Procedimiento Operativo Estándar de Selección y Baremación del Talento Humano NewCare' "
                "y en el instrumento técnico 'Baremacion_Personal_Asistencial_NEWCARE_Optimus_Health.xlsx', "
                "los cuales forman parte integral del Sistema de Gestión de Calidad de Optimus Health IPS y reposan como Anexos Técnicos del presente manual."
            )
        elif txt == "funciones completas":
            p.text = (
                "El equipo de atención se encuentra articulado bajo el principio de liderazgo colaborativo y horizontal promovido por la Fundación New Health. "
                "El Director Médico y Líder del Programa, Dr. Julio Herrera, asume la conducción técnica y estratégica, garantizando la deliberación clínica colegiada, "
                "la supervisión de los planes de atención individualizados (PAI) y la adecuada transición de los usuarios entre los niveles asistenciales. "
                "La gestión de casos es liderada por la Jefatura de Enfermería y la Coordinación Asistencial, quienes realizan el monitoreo activo de la cohorte."
            )
        elif txt == "perfiles completos como anexos, req minimos, criterios de la baremacion":
            p.text = (
                "Los perfiles de cargo, requisitos mínimos habilitantes (según Resolución 3100 de 2019 de Minsalud) y competencias deseables en cronicidad compleja "
                "y cuidados paliativos para cada integrante del organigrama se encuentran formalizados y adoptados. Cada profesional cuenta con evaluación de competencias "
                "y un plan individual de inducción y capacitación continua orientado al cumplimiento de los estándares de certificación NewCare."
            )
        elif txt == "Incluir nombres":
            p.text = (
                "La asignación nominal del equipo tratante se mantiene actualizada en el sistema de gestión de talento humano de Optimus Health. "
                "Cada profesional cuenta con asignación de jornada, agenda asistencial delimitada por modalidad (intramural y extramural/domiciliaria) "
                "y responsabilidades específicas en la Junta de Cronicidad Clínica, el Comité de Calidad y la Asamblea de Pacientes."
            )

    # Update Table 4 if any rows are missing names or details
    t4 = doc.tables[4]
    for row in t4.rows:
        cargo = row.cells[0].text.strip()
        nombre = row.cells[3].text.strip()
        if "Médico General" in cargo and not nombre:
            row.cells[3].text = "Equipo Médico Domiciliario e Intramural Optimus Health"
        elif "Personal Administrativo" in cargo and not nombre:
            row.cells[3].text = "Equipo de Admisiones, SIAU y Gestión Documental Optimus Health"

    doc.save(src_file)
    print(f"Manual interno actualizado con éxito: {src_file}")

if __name__ == "__main__":
    update_manual()
