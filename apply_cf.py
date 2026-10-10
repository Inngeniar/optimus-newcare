import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill, Font

def apply_cf():
    wb_path = "Baremacion_Personal_Asistencial_NEWCARE_Optimus_Health.xlsx"
    wb = openpyxl.load_workbook(wb_path)

    green_fill = PatternFill(start_color='E2F0D9', end_color='E2F0D9', fill_type='solid')
    green_font = Font(name='Calibri', size=9, bold=True, color='276A3C')

    yellow_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
    yellow_font = Font(name='Calibri', size=9, bold=True, color='B25900')

    red_fill = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')
    red_font = Font(name='Calibri', size=9, bold=True, color='C00000')

    # Sheet 3
    ws3 = wb['3. Baremo CV (50 Puntos)']
    ws3.conditional_formatting.add('M22:M47', FormulaRule(formula=['ISNUMBER(SEARCH("DESTACADO",M22))'], fill=green_fill, font=green_font))
    ws3.conditional_formatting.add('M22:M47', FormulaRule(formula=['ISNUMBER(SEARCH("IDÓNEO",M22))'], fill=green_fill, font=green_font))
    ws3.conditional_formatting.add('M22:M47', FormulaRule(formula=['ISNUMBER(SEARCH("FORTALECIMIENTO",M22))'], fill=yellow_fill, font=yellow_font))

    ws3.conditional_formatting.add('L7', FormulaRule(formula=['ISNUMBER(SEARCH("DESTACADO",L7))'], fill=green_fill, font=green_font))
    ws3.conditional_formatting.add('L7', FormulaRule(formula=['ISNUMBER(SEARCH("IDÓNEO",L7))'], fill=green_fill, font=green_font))
    ws3.conditional_formatting.add('L7', FormulaRule(formula=['ISNUMBER(SEARCH("FORTALECIMIENTO",L7))'], fill=yellow_fill, font=yellow_font))

    # Sheet 4
    ws4 = wb['4. Psicotecnico y Entrevista']
    ws4.conditional_formatting.add('G26:G51', FormulaRule(formula=['ISNUMBER(SEARCH("APTO",G26))'], fill=green_fill, font=green_font))
    ws4.conditional_formatting.add('G26:G51', FormulaRule(formula=['ISNUMBER(SEARCH("CONDICIONADO",G26))'], fill=yellow_fill, font=yellow_font))

    ws4.conditional_formatting.add('K26:K51', FormulaRule(formula=['ISNUMBER(SEARCH("SOBRESALIENTE",K26))'], fill=green_fill, font=green_font))
    ws4.conditional_formatting.add('K26:K51', FormulaRule(formula=['ISNUMBER(SEARCH("CONDICIONADO",K26))'], fill=yellow_fill, font=yellow_font))
    ws4.conditional_formatting.add('K26:K51', FormulaRule(formula=['ISNUMBER(SEARCH("NO IDÓNEO",K26))'], fill=red_fill, font=red_font))

    wb.save(wb_path)
    print("Formatos condicionales aplicados exitosamente.")

if __name__ == "__main__":
    apply_cf()
