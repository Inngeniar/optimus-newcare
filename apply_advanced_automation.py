import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

def apply_automations():
    wb_path = "Baremacion_Personal_Asistencial_NEWCARE_Optimus_Health.xlsx"
    wb = openpyxl.load_workbook(wb_path)

    # Styles
    c_primary = "073763"
    c_secondary = "1C4587"
    c_accent = "2E75B5"
    c_light = "E8EEF5"
    c_border = "B0C4DE"
    white = "FFFFFF"

    font_hdr = Font(name="Calibri", size=10, bold=True, color=white)
    font_bold = Font(name="Calibri", size=10, bold=True, color="000000")
    font_normal = Font(name="Calibri", size=10, color="000000")
    fill_hdr = PatternFill(start_color=c_primary, end_color=c_primary, fill_type="solid")
    fill_sec = PatternFill(start_color=c_secondary, end_color=c_secondary, fill_type="solid")
    fill_light = PatternFill(start_color=c_light, end_color=c_light, fill_type="solid")
    thin_border = Side(border_style="thin", color=c_border)
    border_cell = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)

    # =========================================================================
    # 1. AUTOMATIZACIÓN EN SHEET 2 (FILTROS EXCLUYENTES SI-NO)
    # =========================================================================
    ws2 = wb["2. Filtros Excluyentes SI-NO"]

    # Validación de datos para resultados CUMPLE / NO CUMPLE
    dv_filtros_res = DataValidation(type="list", formula1='"CUMPLE,NO CUMPLE"', allow_blank=False)
    ws2.add_data_validation(dv_filtros_res)
    dv_filtros_res.add("G5:G13")

    # Fila de dictamen global al final de los filtros
    row_flt_tot = 15
    ws2.merge_cells(f"A{row_flt_tot}:F{row_flt_tot}")
    c_tit = ws2[f"A{row_flt_tot}"]
    c_tit.value = "DICTAMEN GLOBAL DE HABILITACIÓN NORMATIVA Y FILTROS (GATEKEEPER)"
    c_tit.font = font_hdr
    c_tit.fill = fill_hdr
    c_tit.alignment = Alignment(horizontal="center", vertical="center")

    c_res = ws2[f"G{row_flt_tot}"]
    c_res.value = '=IF(COUNTIF(G5:G13,"NO CUMPLE")>0, "NO CUMPLE (DESCARTADO)", "100% CUMPLE (HABILITADO)")'
    c_res.font = Font(name="Calibri", size=10, bold=True, color="073763")
    c_res.fill = fill_light
    c_res.alignment = Alignment(horizontal="center", vertical="center")

    for col in range(1, 8):
        ws2.cell(row=row_flt_tot, column=col).border = border_cell
    ws2.row_dimensions[row_flt_tot].height = 24

    # =========================================================================
    # 2. AUTOMATIZACIÓN EN SHEET 6 (REGISTRO Y EVALUADOR)
    # =========================================================================
    ws6 = wb["6. Registro y Evaluador"]

    # Validaciones de datos
    cargos_list = (
        '"Director Médico / Líder NewCare,'
        'Coordinadora Asistencial,'
        'Gestor de Casos (Case Manager),'
        'Médico Especialista Internista,'
        'Médico Especialista Intensivista,'
        'Médico Especialista Pediatra,'
        'Médico General Domiciliario / Ambulatorio,'
        'Auxiliar de Enfermería Domiciliaria,'
        'Profesional de Fisioterapia,'
        'Profesional de Fonoaudiología,'
        'Profesional de Psicología,'
        'Profesional de Nutrición,'
        'Profesional de Trabajo Social"'
    )
    dv_cargos = DataValidation(type="list", formula1=cargos_list, allow_blank=True)
    ws6.add_data_validation(dv_cargos)
    dv_cargos.add("C5:C35")

    dv_filtros_tabla = DataValidation(type="list", formula1='"CUMPLE 100%,NO CUMPLE (DESCARTADO)"', allow_blank=True)
    ws6.add_data_validation(dv_filtros_tabla)
    dv_filtros_tabla.add("D5:D35")

    # Actualización inteligente de fórmulas en filas 5 a 35
    for r in range(5, 36):
        # Columna H: Puntaje Final
        # Regla: Si no hay nombre, en blanco. Si filtros <> "CUMPLE 100%", puntaje = 0. Si cumple, suma E+F+G
        formula_h = f'=IF(B{r}="","",IF(D{r}<>"CUMPLE 100%",0,ROUND(E{r}+F{r}+G{r},1)))'
        cell_h = ws6.cell(row=r, column=8, value=formula_h)
        cell_h.font = font_bold
        cell_h.alignment = Alignment(horizontal="center", vertical="center")

        # Columna I: Dictamen Final
        # Regla: Si no hay nombre, en blanco. Si filtros <> "CUMPLE 100%", descarta. Si puntaje >= 85 Sobresaliente, >= 70 Condicionado, sino No Idóneo
        formula_i = (
            f'=IF(B{r}="","",'
            f'IF(D{r}<>"CUMPLE 100%","NO IDÓNEO (DESCARTADO EN FILTROS HABILITANTES)",'
            f'IF(H{r}>=85,"IDÓNEO SOBRESALIENTE",'
            f'IF(H{r}>=70,"IDÓNEO CONDICIONADO","NO IDÓNEO"))))'
        )
        cell_i = ws6.cell(row=r, column=9, value=formula_i)
        cell_i.font = font_bold
        cell_i.alignment = Alignment(horizontal="center", vertical="center")

    # =========================================================================
    # 3. CUADRO DE MANDO EJECUTIVO (KPIs EN VIVO EN COLUMNAS K A M)
    # =========================================================================
    ws6.merge_cells("K4:M4")
    ws6["K4"] = "CUADRO DE MANDO EJECUTIVO (KPIs EN VIVO)"
    ws6["K4"].font = font_hdr
    ws6["K4"].fill = fill_hdr
    ws6["K4"].alignment = Alignment(horizontal="center", vertical="center")
    ws6.row_dimensions[4].height = 28

    kpis = [
        ("Total Candidatos Registrados", '=COUNTA(B5:B35)-COUNTIF(B5:B35,"")', False),
        ("Superan Filtros Habilitantes", '=COUNTIF(D5:D35,"CUMPLE 100%")', False),
        ("Idóneos Sobresalientes (>=85)", '=COUNTIF(I5:I35,"IDÓNEO SOBRESALIENTE")', False),
        ("Idóneos Condicionados (70-84)", '=COUNTIF(I5:I35,"IDÓNEO CONDICIONADO")', False),
        ("Descartados / No Idóneos", '=COUNTIF(I5:I35,"*NO*")', False),
        ("Promedio de Puntaje (Aprobados)", '=IFERROR(ROUND(AVERAGEIF(H5:H35,">0"),1),0)', False),
        ("Tasa de Aprobación Global", '=IFERROR(ROUND((M7+M8)/MAX(1,M5),3),0)', True)
    ]

    for idx, (kpi_title, kpi_formula, is_pct) in enumerate(kpis, start=5):
        ws6.merge_cells(f"K{idx}:L{idx}")
        c_k = ws6[f"K{idx}"]
        c_k.value = kpi_title
        c_k.font = font_bold if "Tasa" in kpi_title or "Total" in kpi_title else font_normal
        c_k.alignment = Alignment(horizontal="left", vertical="center")

        c_v = ws6[f"M{idx}"]
        c_v.value = kpi_formula
        c_v.font = font_bold
        c_v.alignment = Alignment(horizontal="center", vertical="center")
        if is_pct:
            c_v.number_format = '0.0%'
        else:
            c_v.number_format = '#,##0.0' if "Promedio" in kpi_title else '#,##0'

        for col in range(11, 14):
            c_cell = ws6.cell(row=idx, column=col)
            c_cell.border = border_cell
            if idx % 2 == 0:
                c_cell.fill = PatternFill(start_color="F7F9FC", end_color="F7F9FC", fill_type="solid")
            if "Tasa" in kpi_title:
                c_cell.fill = fill_light

    ws6.column_dimensions["K"].width = 24
    ws6.column_dimensions["L"].width = 12
    ws6.column_dimensions["M"].width = 16

    # =========================================================================
    # 4. FORMATO CONDICIONAL DINÁMICO EN COLUMNA I (DICTAMEN)
    # =========================================================================
    green_fill = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")
    green_font = Font(name="Calibri", size=10, bold=True, color="276A3C")
    yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    yellow_font = Font(name="Calibri", size=10, bold=True, color="B25900")
    red_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    red_font = Font(name="Calibri", size=10, bold=True, color="C00000")

    rule_green = FormulaRule(formula=['ISNUMBER(SEARCH("SOBRESALIENTE",I5))'], fill=green_fill, font=green_font)
    rule_yellow = FormulaRule(formula=['ISNUMBER(SEARCH("CONDICIONADO",I5))'], fill=yellow_fill, font=yellow_font)
    rule_red = FormulaRule(formula=['OR(ISNUMBER(SEARCH("NO IDÓNEO",I5)),ISNUMBER(SEARCH("DESCARTADO",I5)))'], fill=red_fill, font=red_font)

    ws6.conditional_formatting.add("I5:I35", rule_green)
    ws6.conditional_formatting.add("I5:I35", rule_yellow)
    ws6.conditional_formatting.add("I5:I35", rule_red)

    wb.save(wb_path)
    print(f"Automatizaciones aplicadas exitosamente sobre: {wb_path}")

if __name__ == "__main__":
    apply_automations()
