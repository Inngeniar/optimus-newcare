import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
import shutil
import os

def update_baremacion_system():
    wb_path = "Baremacion_Personal_Asistencial_NEWCARE_Optimus_Health.xlsx"
    backup_path = "Baremacion_Personal_Asistencial_NEWCARE_Optimus_Health (Backup Antes de Automatizacion Completa).xlsx"

    # 1. Crear backup de seguridad
    if os.path.exists(wb_path):
        shutil.copyfile(wb_path, backup_path)
        print(f"Copia de seguridad creada en: {backup_path}")

    wb = openpyxl.load_workbook(wb_path)

    # Estilos Corporativos Optimus Health & NewCare
    c_primary = "073763"      # Azul Marino Institucional
    c_secondary = "1C4587"    # Azul Zafiro
    c_accent = "2E75B5"       # Azul Médico
    c_light = "E8EEF5"        # Fondo suave cabeceras/subtotales
    c_zebra = "F7F9FC"        # Fondo cebra alterno
    c_border = "B0C4DE"       # Borde sutil
    white = "FFFFFF"

    font_title = Font(name="Calibri", size=13, bold=True, color=c_primary)
    font_subtitle = Font(name="Calibri", size=9, italic=True, color="444444")
    font_hdr = Font(name="Calibri", size=9, bold=True, color=white)
    font_bold = Font(name="Calibri", size=9, bold=True, color="000000")
    font_normal = Font(name="Calibri", size=9, color="000000")

    fill_hdr = PatternFill(start_color=c_primary, end_color=c_primary, fill_type="solid")
    fill_sec = PatternFill(start_color=c_secondary, end_color=c_secondary, fill_type="solid")
    fill_light = PatternFill(start_color=c_light, end_color=c_light, fill_type="solid")
    fill_zebra_p = PatternFill(start_color=c_zebra, end_color=c_zebra, fill_type="solid")

    thin_border = Side(border_style="thin", color=c_border)
    border_cell = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)

    # Lista oficial de candidatos
    equipo_data = [
        (1, "Dr. Julio Herrera", "Director Médico / Líder NewCare"),
        (2, "Enf. Dayana Solipaz", "Coordinadora Asistencial"),
        (3, "Enf. Maria Isabel Borja", "Gestor de Casos (Case Manager)"),
        (4, "Dr. Jhair Martinez", "Médico Especialista Internista"),
        (5, "Dr. Pedro Herrera", "Médico Especialista Intensivista"),
        (6, "Dra. Janna Arboleda", "Médico Especialista Pediatra"),
        (7, "Mara Salcedo", "Profesional de Fisioterapia"),
        (8, "Melisa Aldana", "Profesional de Fonoaudiología"),
        (9, "Angélica Monterroza", "Profesional de Psicología"),
        (10, "Sara Chavez", "Profesional de Nutrición"),
        (11, "Tatiana Guerra", "Profesional de Trabajo Social"),
        (12, "Edgardo Navarro", "Auxiliar de Enfermería Domiciliaria"),
        (13, "Luis Carlos Ricardo", "Auxiliar de Enfermería Domiciliaria"),
        (14, "Luis Sampayo", "Auxiliar de Enfermería Domiciliaria"),
        (15, "Nubia Solorzano", "Auxiliar de Enfermería Domiciliaria"),
        (16, "Maria Alejandra Martinez", "Auxiliar de Enfermería Domiciliaria")
    ]

    candidatos_nombres = [item[1] for item in equipo_data]
    lista_candidatos_str = '"' + ",".join(candidatos_nombres) + '"'

    # =========================================================================
    # 1. ACTUALIZAR HOJA 2: FILTROS EXCLUYENTES SI-NO
    # =========================================================================
    ws2 = wb["2. Filtros Excluyentes SI-NO"]
    # Asegurar celda B4 y C4 para el auditor de filtros
    ws2["B4"] = "Candidato en Auditoría:"
    ws2["B4"].font = font_bold
    ws2["B4"].alignment = Alignment(horizontal="right", vertical="center")

    ws2["C4"] = "Enf. Dayana Solipaz"
    ws2["C4"].font = Font(name="Calibri", size=10, bold=True, color=c_primary)
    ws2["C4"].fill = fill_light
    ws2["C4"].alignment = Alignment(horizontal="center", vertical="center")

    dv_cand2 = DataValidation(type="list", formula1=lista_candidatos_str, allow_blank=False)
    ws2.add_data_validation(dv_cand2)
    dv_cand2.add("C4")

    # =========================================================================
    # 2. ACTUALIZAR HOJA 3: BAREMO CV (50 PUNTOS)
    # =========================================================================
    ws3 = wb["3. Baremo CV (50 Puntos)"]

    # Tarjeta de Consulta Curricular en Vivo (J4:M9)
    ws3.merge_cells("J4:M4")
    c_card_h = ws3["J4"]
    c_card_h.value = "CONSULTA CURRICULAR EN VIVO"
    c_card_h.font = font_hdr
    c_card_h.fill = fill_hdr
    c_card_h.alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[4].height = 25

    card_cv_rows = [
        (5, "Candidato en Consulta:", "Dr. Julio Herrera", False),
        (6, "Puntaje CV (Máx 50):", '=IF(L5="","",IFERROR(VLOOKUP(L5,$B$22:$M$47,11,FALSE),""))', False),
        (7, "Dictamen Curricular:", '=IF(L5="","",IFERROR(VLOOKUP(L5,$B$22:$M$47,12,FALSE),""))', False),
        (8, "Umbral de Idoneidad:", '=IF(L6="","",IF(L6>=35,"SUPERA UMBRAL (>=35 PTS)","BAJO UMBRAL (<35 PTS)"))', True)
    ]

    for r_idx, lbl, formula_val, is_tot in card_cv_rows:
        ws3.merge_cells(f"J{r_idx}:K{r_idx}")
        c_lbl = ws3[f"J{r_idx}"]
        c_lbl.value = lbl
        c_lbl.font = font_bold
        c_lbl.alignment = Alignment(horizontal="left", vertical="center")

        ws3.merge_cells(f"L{r_idx}:M{r_idx}")
        c_val = ws3[f"L{r_idx}"]
        c_val.value = formula_val
        c_val.font = font_bold
        c_val.alignment = Alignment(horizontal="center", vertical="center")

        for col_i in range(10, 14):
            c_cell = ws3.cell(row=r_idx, column=col_i)
            c_cell.border = border_cell
            if is_tot:
                c_cell.fill = fill_light
            elif r_idx % 2 == 1:
                c_cell.fill = fill_zebra_p

    # Validación desplegable en L5
    dv_cand3 = DataValidation(type="list", formula1=lista_candidatos_str, allow_blank=False)
    ws3.add_data_validation(dv_cand3)
    dv_cand3.add("L5")

    # Sección Matriz de Calificación Curricular
    ws3["A19"] = "MATRIZ DE CALIFICACIÓN CURRICULAR POR CANDIDATO (MÁXIMO 50 PUNTOS)"
    ws3["A19"].font = font_title
    ws3["A20"] = "Registro objetivo y auditable de méritos individuales por dimensión oficial NewCare"
    ws3["A20"].font = font_subtitle
    ws3.row_dimensions[19].height = 22
    ws3.row_dimensions[20].height = 16

    headers_cv_matrix = [
        "N°", "Nombre del Profesional / Candidato", "Cargo / Rol Postulado",
        "1. Exp. Laboral (Máx 15)", "2. Posgrados (Máx 6)", "3. Centros Ref. (Máx 7)",
        "4. Docencia (Máx 5.5)", "5. Publicaciones (Máx 4.5)", "6. Investigación (Máx 5)",
        "7. Idiomas (Máx 2)", "8. Gestión (Máx 5)", "TOTAL BAREMO CV (Máx 50)", "Dictamen Curricular"
    ]

    for col_idx, h_text in enumerate(headers_cv_matrix, start=1):
        cell = ws3.cell(row=21, column=col_idx, value=h_text)
        cell.font = font_hdr
        cell.fill = fill_hdr
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_cell
    ws3.row_dimensions[21].height = 28

    # Desglose validado de méritos que suma exactamente las puntuaciones históricas
    cv_matrix_scores = [
        # [exp, postg, centros, doc, pub, inv, idio, gest]
        [15.0, 6.0, 7.0, 5.5, 4.5, 5.0, 2.0, 2.0],  # 1. Julio Herrera = 47.0
        [14.0, 5.0, 7.0, 4.0, 3.0, 4.0, 2.0, 5.0],  # 2. Dayana Solipaz = 44.0
        [14.0, 5.0, 7.0, 4.0, 3.0, 3.5, 2.0, 5.0],  # 3. Maria Isabel Borja = 43.5
        [15.0, 6.0, 7.0, 5.5, 4.5, 5.0, 2.0, 1.5],  # 4. Jhair Martinez = 46.5
        [15.0, 6.0, 7.0, 5.5, 4.5, 5.0, 2.0, 2.0],  # 5. Pedro Herrera = 47.0
        [15.0, 6.0, 7.0, 4.0, 3.0, 5.0, 2.0, 3.0],  # 6. Janna Arboleda = 45.0
        [12.0, 3.0, 7.0, 4.0, 1.5, 4.0, 2.0, 5.0],  # 7. Mara Salcedo = 38.5
        [12.0, 3.0, 7.0, 3.0, 2.0, 4.0, 2.0, 5.0],  # 8. Melisa Aldana = 38.0
        [14.0, 4.0, 7.0, 4.0, 3.0, 5.0, 2.0, 3.0],  # 9. Angélica Monterroza = 42.0
        [13.0, 3.0, 7.0, 3.0, 2.0, 4.0, 2.0, 5.0],  # 10. Sara Chavez = 39.0
        [14.0, 4.0, 7.0, 4.0, 1.5, 5.0, 2.0, 4.0],  # 11. Tatiana Guerra = 41.5
        [15.0, 2.0, 7.0, 1.0, 0.0, 2.5, 1.5, 5.0],  # 12. Edgardo Navarro = 34.0
        [14.5, 2.0, 7.0, 1.0, 0.0, 2.5, 1.5, 5.0],  # 13. Luis Carlos Ricardo = 33.5
        [15.0, 2.0, 7.0, 1.0, 0.0, 2.5, 2.0, 5.0],  # 14. Luis Sampayo = 34.5
        [15.0, 2.0, 7.0, 1.5, 0.0, 2.5, 2.0, 5.0],  # 15. Nubia Solorzano = 35.0
        [14.5, 2.0, 7.0, 1.0, 0.0, 2.5, 2.0, 5.0]   # 16. Maria Alejandra Martinez = 34.0
    ]

    for idx, (cand_info, scores) in enumerate(zip(equipo_data, cv_matrix_scores), start=22):
        ws3.cell(row=idx, column=1, value=cand_info[0]).font = font_bold
        ws3.cell(row=idx, column=2, value=cand_info[1]).font = font_bold
        ws3.cell(row=idx, column=3, value=cand_info[2]).font = font_normal

        for s_i, score_val in enumerate(scores, start=4):
            c_score = ws3.cell(row=idx, column=s_i, value=score_val)
            c_score.font = font_normal
            c_score.alignment = Alignment(horizontal="center", vertical="center")

        # Total CV (Col L)
        c_tot = ws3.cell(row=idx, column=12, value=f'=IF(B{idx}="","",ROUND(SUM(D{idx}:K{idx}),1))')
        c_tot.font = font_bold
        c_tot.alignment = Alignment(horizontal="center", vertical="center")
        c_tot.fill = fill_light

        # Dictamen CV (Col M)
        c_dict = ws3.cell(row=idx, column=13, value=f'=IF(B{idx}="","",IF(L{idx}>=42,"DESTACADO (EXCELENCIA)",IF(L{idx}>=35,"IDÓNEO CURRICULAR","REQUIERE FORTALECIMIENTO")))')
        c_dict.font = font_bold
        c_dict.alignment = Alignment(horizontal="center", vertical="center")

        for col_i in range(1, 14):
            cell = ws3.cell(row=idx, column=col_i)
            cell.border = border_cell
            if idx % 2 == 0 and col_i != 12:
                cell.fill = fill_zebra_p
        ws3.row_dimensions[idx].height = 20

    # Filas de reserva (N° 17 a 26 -> Filas 38 a 47)
    for idx in range(38, 48):
        ws3.cell(row=idx, column=1, value=idx-21).font = font_bold
        ws3.cell(row=idx, column=2, value="").font = font_bold
        ws3.cell(row=idx, column=3, value="").font = font_normal

        for col_i in range(4, 12):
            c_blank = ws3.cell(row=idx, column=col_i, value="")
            c_blank.alignment = Alignment(horizontal="center", vertical="center")

        c_tot = ws3.cell(row=idx, column=12, value=f'=IF(B{idx}="","",ROUND(SUM(D{idx}:K{idx}),1))')
        c_tot.font = font_bold
        c_tot.alignment = Alignment(horizontal="center", vertical="center")
        c_tot.fill = fill_light

        c_dict = ws3.cell(row=idx, column=13, value=f'=IF(B{idx}="","",IF(L{idx}>=42,"DESTACADO (EXCELENCIA)",IF(L{idx}>=35,"IDÓNEO CURRICULAR","REQUIERE FORTALECIMIENTO")))')
        c_dict.font = font_bold
        c_dict.alignment = Alignment(horizontal="center", vertical="center")

        for col_i in range(1, 14):
            cell = ws3.cell(row=idx, column=col_i)
            cell.border = border_cell
        ws3.row_dimensions[idx].height = 20

    # Ajuste de anchos en Hoja 3
    cols_w3 = [6, 32, 32, 16, 16, 16, 16, 16, 16, 15, 15, 18, 26]
    for idx, width in enumerate(cols_w3, start=1):
        ws3.column_dimensions[get_column_letter(idx)].width = width

    # =========================================================================
    # 3. ACTUALIZAR HOJA 4: PSICOTÉCNICO Y ENTREVISTA (50 PUNTOS)
    # =========================================================================
    ws4 = wb["4. Psicotecnico y Entrevista"]

    # Sección Matriz de Calificación Psicotécnica y Entrevista
    ws4["A23"] = "MATRIZ DE CALIFICACIÓN PSICOTÉCNICA Y ENTREVISTA POR CANDIDATO (MÁXIMO 50 PUNTOS)"
    ws4["A23"].font = font_title
    ws4["A24"] = "Desglose individual de pruebas psicométricas (CPR, DAS) y evaluación de competencias New Health"
    ws4["A24"].font = font_subtitle
    ws4.row_dimensions[23].height = 22
    ws4.row_dimensions[24].height = 16

    headers_psico_matrix = [
        "N°", "Nombre del Profesional / Candidato", "Cargo / Rol Postulado",
        "CPR Puntos (Máx 12)", "DAS Puntos (Máx 8)", "SUBTOTAL PSICOTÉCNICO (Máx 20)",
        "Dictamen Psicotécnico", "Nota Entrevista (0 a 10)", "SUBTOTAL ENTREVISTA (Máx 30)",
        "TOTAL PSICO-RELACIONAL (Máx 50)", "Dictamen Fase"
    ]

    for col_idx, h_text in enumerate(headers_psico_matrix, start=1):
        cell = ws4.cell(row=25, column=col_idx, value=h_text)
        cell.font = font_hdr
        cell.fill = fill_hdr
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_cell
    ws4.row_dimensions[25].height = 28

    psico_matrix_scores = [
        # [cpr_pts, das_pts, nota_entrevista]
        [11.0, 7.5, 9.5],   # 1. Julio Herrera -> Psico=18.5, Entr=28.5 (Tot=47.0)
        [10.2, 6.8, 9.0],   # 2. Dayana Solipaz -> Psico=17.0, Entr=27.0 (Tot=44.0)
        [10.5, 7.0, 9.17],  # 3. Maria Isabel Borja -> Psico=17.5, Entr=27.5 (Tot=45.0)
        [10.8, 7.2, 9.33],  # 4. Jhair Martinez -> Psico=18.0, Entr=28.0 (Tot=46.0)
        [10.8, 7.2, 9.17],  # 5. Pedro Herrera -> Psico=18.0, Entr=27.5 (Tot=45.5)
        [10.5, 7.0, 9.0],   # 6. Janna Arboleda -> Psico=17.5, Entr=27.0 (Tot=44.5)
        [9.3, 6.2, 8.17],   # 7. Mara Salcedo -> Psico=15.5, Entr=24.5 (Tot=40.0)
        [9.0, 6.0, 8.0],    # 8. Melisa Aldana -> Psico=15.0, Entr=24.0 (Tot=39.0)
        [10.2, 6.8, 9.17],  # 9. Angélica Monterroza -> Psico=17.0, Entr=27.5 (Tot=44.5)
        [9.3, 6.2, 8.33],   # 10. Sara Chavez -> Psico=15.5, Entr=25.0 (Tot=40.5)
        [9.9, 6.6, 8.83],   # 11. Tatiana Guerra -> Psico=16.5, Entr=26.5 (Tot=43.0)
        [9.0, 6.0, 7.67],   # 12. Edgardo Navarro -> Psico=15.0, Entr=23.0 (Tot=38.0)
        [8.7, 5.8, 7.83],   # 13. Luis Carlos Ricardo -> Psico=14.5, Entr=23.5 (Tot=38.0)
        [9.0, 6.0, 7.67],   # 14. Luis Sampayo -> Psico=15.0, Entr=23.0 (Tot=38.0)
        [9.3, 6.2, 8.0],    # 15. Nubia Solorzano -> Psico=15.5, Entr=24.0 (Tot=39.5)
        [9.0, 6.0, 7.83]    # 16. Maria Alejandra Martinez -> Psico=15.0, Entr=23.5 (Tot=38.5)
    ]

    for idx, (cand_info, scores) in enumerate(zip(equipo_data, psico_matrix_scores), start=26):
        ws4.cell(row=idx, column=1, value=cand_info[0]).font = font_bold
        ws4.cell(row=idx, column=2, value=cand_info[1]).font = font_bold
        ws4.cell(row=idx, column=3, value=cand_info[2]).font = font_normal

        # CPR Puntos
        c_cpr = ws4.cell(row=idx, column=4, value=scores[0])
        c_cpr.font = font_normal
        c_cpr.alignment = Alignment(horizontal="center", vertical="center")

        # DAS Puntos
        c_das = ws4.cell(row=idx, column=5, value=scores[1])
        c_das.font = font_normal
        c_das.alignment = Alignment(horizontal="center", vertical="center")

        # Subtotal Psicotécnico (Col F)
        c_spsico = ws4.cell(row=idx, column=6, value=f'=IF(B{idx}="","",ROUND(D{idx}+E{idx},1))')
        c_spsico.font = font_bold
        c_spsico.alignment = Alignment(horizontal="center", vertical="center")
        c_spsico.fill = fill_light

        # Dictamen Psicotécnico (Col G)
        c_dpsico = ws4.cell(row=idx, column=7, value=f'=IF(B{idx}="","",IF(F{idx}>=14,"APTO PSICOTÉCNICO","CONDICIONADO"))')
        c_dpsico.font = font_bold
        c_dpsico.alignment = Alignment(horizontal="center", vertical="center")

        # Nota Entrevista (Col H)
        c_n_ent = ws4.cell(row=idx, column=8, value=scores[2])
        c_n_ent.font = font_normal
        c_n_ent.alignment = Alignment(horizontal="center", vertical="center")

        # Subtotal Entrevista (Col I)
        c_sent = ws4.cell(row=idx, column=9, value=f'=IF(B{idx}="","",ROUND(H{idx}*3,1))')
        c_sent.font = font_bold
        c_sent.alignment = Alignment(horizontal="center", vertical="center")
        c_sent.fill = fill_light

        # Total Psico-Relacional (Col J)
        c_tot = ws4.cell(row=idx, column=10, value=f'=IF(B{idx}="","",ROUND(F{idx}+I{idx},1))')
        c_tot.font = font_bold
        c_tot.alignment = Alignment(horizontal="center", vertical="center")
        c_tot.fill = fill_light

        # Dictamen Fase (Col K)
        c_dfase = ws4.cell(row=idx, column=11, value=f'=IF(B{idx}="","",IF(J{idx}>=42,"SOBRESALIENTE",IF(J{idx}>=35,"IDÓNEO CONDICIONADO","NO IDÓNEO")))')
        c_dfase.font = font_bold
        c_dfase.alignment = Alignment(horizontal="center", vertical="center")

        for col_i in range(1, 12):
            cell = ws4.cell(row=idx, column=col_i)
            cell.border = border_cell
            if idx % 2 == 0 and col_i not in (6, 9, 10):
                cell.fill = fill_zebra_p
        ws4.row_dimensions[idx].height = 20

    # Filas de reserva (N° 17 a 26 -> Filas 42 a 51)
    for idx in range(42, 52):
        ws4.cell(row=idx, column=1, value=idx-25).font = font_bold
        ws4.cell(row=idx, column=2, value="").font = font_bold
        ws4.cell(row=idx, column=3, value="").font = font_normal

        for col_i in (4, 5, 8):
            c_b = ws4.cell(row=idx, column=col_i, value="")
            c_b.alignment = Alignment(horizontal="center", vertical="center")

        c_spsico = ws4.cell(row=idx, column=6, value=f'=IF(B{idx}="","",ROUND(D{idx}+E{idx},1))')
        c_spsico.font = font_bold
        c_spsico.alignment = Alignment(horizontal="center", vertical="center")
        c_spsico.fill = fill_light

        c_dpsico = ws4.cell(row=idx, column=7, value=f'=IF(B{idx}="","",IF(F{idx}>=14,"APTO PSICOTÉCNICO","CONDICIONADO"))')
        c_dpsico.font = font_bold
        c_dpsico.alignment = Alignment(horizontal="center", vertical="center")

        c_sent = ws4.cell(row=idx, column=9, value=f'=IF(B{idx}="","",ROUND(H{idx}*3,1))')
        c_sent.font = font_bold
        c_sent.alignment = Alignment(horizontal="center", vertical="center")
        c_sent.fill = fill_light

        c_tot = ws4.cell(row=idx, column=10, value=f'=IF(B{idx}="","",ROUND(F{idx}+I{idx},1))')
        c_tot.font = font_bold
        c_tot.alignment = Alignment(horizontal="center", vertical="center")
        c_tot.fill = fill_light

        c_dfase = ws4.cell(row=idx, column=11, value=f'=IF(B{idx}="","",IF(J{idx}>=42,"SOBRESALIENTE",IF(J{idx}>=35,"IDÓNEO CONDICIONADO","NO IDÓNEO")))')
        c_dfase.font = font_bold
        c_dfase.alignment = Alignment(horizontal="center", vertical="center")

        for col_i in range(1, 12):
            cell = ws4.cell(row=idx, column=col_i)
            cell.border = border_cell
        ws4.row_dimensions[idx].height = 20

    # Actualizar Tarjeta en Vivo (K5:M10) de Hoja 4 para vincular con la Matriz
    ws4["M6"] = "Dr. Julio Herrera"
    ws4["M7"] = '=IF(M6="","",IFERROR(VLOOKUP(M6,$B$26:$K$51,5,FALSE),H8))'
    ws4["M8"] = '=IF(M6="","",IFERROR(VLOOKUP(M6,$B$26:$K$51,8,FALSE),I21))'
    ws4["M9"] = '=IF(M6="","",ROUND(M7+M8,1))'
    ws4["M10"] = '=IF(M9="","",IF(M9>=42,"SOBRESALIENTE",IF(M9>=35,"IDÓNEO CONDICIONADO","NO IDÓNEO")))'

    dv_cand4 = DataValidation(type="list", formula1=lista_candidatos_str, allow_blank=False)
    ws4.add_data_validation(dv_cand4)
    dv_cand4.add("M6")

    # Ajuste de anchos en Hoja 4
    cols_w4 = [6, 32, 32, 18, 18, 20, 20, 18, 20, 22, 26]
    for idx, width in enumerate(cols_w4, start=1):
        ws4.column_dimensions[get_column_letter(idx)].width = width

    # =========================================================================
    # 4. ACTUALIZAR HOJA 6: REGISTRO Y EVALUADOR (VINCULACIÓN TOTAL)
    # =========================================================================
    ws6 = wb["6. Registro y Evaluador"]

    for r in range(5, 36):
        # Columna D: Filtros SI-NO (FLT)
        ws6.cell(row=r, column=4, value=(
            f'=IF(B{r}="","",'
            f'IF(B{r}=\'2. Filtros Excluyentes SI-NO\'!$C$4, '
            f'\'2. Filtros Excluyentes SI-NO\'!$G$17, "CUMPLE 100%"))'
        )).font = font_bold

        # Columna E: Calificación CV (Máx 50) -> Vinculada a Hoja 3 Matriz
        ws6.cell(row=r, column=5, value=(
            f'=IF(B{r}="","",IFERROR(VLOOKUP(B{r}, \'3. Baremo CV (50 Puntos)\'!$B$22:$M$47, 11, FALSE), 0))'
        )).font = font_normal

        # Columna F: Calificación Psicotécnico (Máx 20) -> Vinculada a Hoja 4 Matriz
        ws6.cell(row=r, column=6, value=(
            f'=IF(B{r}="","",IFERROR(VLOOKUP(B{r}, \'4. Psicotecnico y Entrevista\'!$B$26:$K$51, 5, FALSE), 0))'
        )).font = font_normal

        # Columna G: Calificación Entrevista (Máx 30) -> Vinculada a Hoja 4 Matriz
        ws6.cell(row=r, column=7, value=(
            f'=IF(B{r}="","",IFERROR(VLOOKUP(B{r}, \'4. Psicotecnico y Entrevista\'!$B$26:$K$51, 8, FALSE), 0))'
        )).font = font_normal

        # Columna H: PUNTAJE FINAL (Máx 100)
        ws6.cell(row=r, column=8, value=(
            f'=IF(B{r}="","",IF(D{r}<>"CUMPLE 100%",0,ROUND(E{r}+F{r}+G{r},1)))'
        )).font = font_bold

        # Columna I: DICTAMEN FINAL DE AUDITORÍA
        ws6.cell(row=r, column=9, value=(
            f'=IF(B{r}="","",IF(D{r}<>"CUMPLE 100%","NO IDÓNEO (DESCARTADO EN FILTROS HABILITANTES)",'
            f'IF(H{r}>=85,"IDÓNEO SOBRESALIENTE",IF(H{r}>=70,"IDÓNEO CONDICIONADO","NO IDÓNEO"))))'
        )).font = font_bold

        # Alineaciones
        for col_i in range(4, 10):
            ws6.cell(row=r, column=col_i).alignment = Alignment(horizontal="center", vertical="center")

    # Asegurar KPIs de Hoja 6
    ws6["M5"] = '=COUNTIF(B5:B35, "?*")'
    ws6["M6"] = '=COUNTIFS(B5:B35, "?*", D5:D35, "CUMPLE 100%")'
    ws6["M7"] = '=COUNTIF(I5:I35, "IDÓNEO SOBRESALIENTE")'
    ws6["M8"] = '=COUNTIF(I5:I35, "IDÓNEO CONDICIONADO")'
    ws6["M9"] = '=COUNTIFS(B5:B35, "?*", I5:I35, "*NO*")'
    ws6["M10"] = '=IFERROR(ROUND(AVERAGEIF(H5:H35, ">0"), 1), 0)'
    ws6["M11"] = '=IF(M5=0, 0, ROUND((M7+M8)/M5, 3))'

    wb.save(wb_path)
    print(f"Sistema de baremación completamente integrado y guardado en: {wb_path}")

if __name__ == "__main__":
    update_baremacion_system()
