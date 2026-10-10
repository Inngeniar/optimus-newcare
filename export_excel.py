"""
export_excel.py
Generador avanzado de informes en Excel (.xlsx) para el Análisis GAP NewCare de Optimus Health IPS.
Produce un libro de trabajo profesional, enriquecido con estilos corporativos, semáforos,
resumen ejecutivo, matriz detallada y plan de mejora continua.
"""

import io
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_excel_report(evaluation: Dict[str, Any], items: List[Dict[str, Any]]) -> bytes:
    """
    Genera el archivo Excel en memoria (bytes) para descarga directa.
    """
    wb = openpyxl.Workbook()
    # Eliminar hoja por defecto
    wb.remove(wb.active)
    
    # ----------------------------------------------------
    # ESTILOS CORPORATIVOS OPTIMUS HEALTH
    # ----------------------------------------------------
    font_family = "Segoe UI"
    
    title_font = Font(name=font_family, size=16, bold=True, color="FFFFFF")
    subtitle_font = Font(name=font_family, size=11, italic=True, color="E2E8F0")
    section_font = Font(name=font_family, size=12, bold=True, color="1E293B")
    header_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    bold_font = Font(name=font_family, size=10, bold=True, color="1E293B")
    regular_font = Font(name=font_family, size=10, color="334155")
    small_font = Font(name=font_family, size=9, color="64748B")
    
    # Rellenos
    fill_navy = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_teal = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    fill_header = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    
    fill_green = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")  # soft green
    font_green = Font(name=font_family, size=10, bold=True, color="065F46")
    fill_yellow = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")  # soft amber
    font_yellow = Font(name=font_family, size=10, bold=True, color="92400E")
    fill_red = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")  # soft red
    font_red = Font(name=font_family, size=10, bold=True, color="991B1B")
    fill_gray = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    font_gray = Font(name=font_family, size=10, italic=True, color="64748B")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center")

    # ====================================================
    # HOJA 1: RESUMEN EJECUTIVO (DASHBOARD)
    # ====================================================
    ws1 = wb.create_sheet(title="Resumen Ejecutivo")
    ws1.views.sheetView[0].showGridLines = True
    
    # Banner Header
    ws1.merge_cells("A1:K2")
    ws1["A1"] = "OPTIMUS HEALTH IPS  |  ANÁLISIS DE BRECHAS (GAP ANALYSIS) PROGRAMA NEWCARE"
    ws1["A1"].font = title_font
    ws1["A1"].fill = fill_teal
    ws1["A1"].alignment = align_center
    
    ws1.merge_cells("A3:K3")
    ws1["A3"] = "Preparación y Diagnóstico de Certificación frente a la Metodología NEWCARE (Fundación New Health)"
    ws1["A3"].font = subtitle_font
    ws1["A3"].fill = fill_navy
    ws1["A3"].alignment = align_center
    
    # KPI Box
    ws1["B5"] = "DIAGNÓSTICO GLOBAL DE CERTIFICACIÓN"
    ws1["B5"].font = section_font
    
    ws1["B6"] = "Cumplimiento Ponderado Estratégico:"
    ws1["C6"] = f"{evaluation['global_compliance']}%"
    ws1["C6"].font = Font(name=font_family, size=12, bold=True, color="0F766E")
    
    ws1["E6"] = "Meta de Certificación:"
    ws1["F6"] = f"{evaluation['target_pct']}%"
    ws1["F6"].font = Font(name=font_family, size=11, bold=True, color="334155")
    
    ws1["H6"] = "Brecha Global (GAP):"
    ws1["I6"] = f"{evaluation['global_gap']}%"
    ws1["I6"].font = Font(name=font_family, size=11, bold=True, color="DC2626")
    
    ws1["B7"] = "Veredicto de Preparación:"
    ws1.merge_cells("C7:F7")
    ws1["C7"] = evaluation["verdict_title"]
    ws1["C7"].font = Font(name=font_family, size=11, bold=True, color="0F766E")
    
    ws1["H7"] = "Distribución de Requisitos:"
    ws1["I7"] = f"🟢 {evaluation['total_completo']}  |  🟡 {evaluation['total_parcial']}  |  🔴 {evaluation['total_no_cumple']}"
    ws1["I7"].font = bold_font

    # Tabla de Resultados por Grupo
    ws1["A9"] = "RESULTADOS POR CAPÍTULO NORMATIVO (13 GRUPOS)"
    ws1["A9"].font = section_font
    
    headers_g = [
        "Cód", "Capítulo Normativo NewCare", "Peso %", "Total Req.", "Cumple Completo (2)",
        "Cumple Parcial (1)", "No Cumple (0)", "No Aplica (NA)", "Puntos Obt.", "Puntos Máx.",
        "% Cumplimiento", "Brecha GAP %", "Semáforo / Estado"
    ]
    
    for col_idx, h in enumerate(headers_g, 1):
        cell = ws1.cell(row=10, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border
        
    row_curr = 11
    for g in evaluation["groups"]:
        ws1.cell(row=row_curr, column=1, value=g["code"]).alignment = align_center
        ws1.cell(row=row_curr, column=2, value=g["name"]).alignment = align_left
        ws1.cell(row=row_curr, column=3, value=f"{g['weight_differential']}%").alignment = align_right
        ws1.cell(row=row_curr, column=4, value=g["total_items"]).alignment = align_center
        ws1.cell(row=row_curr, column=5, value=g["count_completo"]).alignment = align_center
        ws1.cell(row=row_curr, column=6, value=g["count_parcial"]).alignment = align_center
        ws1.cell(row=row_curr, column=7, value=g["count_no_cumple"]).alignment = align_center
        ws1.cell(row=row_curr, column=8, value=g["count_na"]).alignment = align_center
        ws1.cell(row=row_curr, column=9, value=g["points_obtained"]).alignment = align_right
        ws1.cell(row=row_curr, column=10, value=g["max_points"]).alignment = align_right
        
        c_pct = ws1.cell(row=row_curr, column=11, value=f"{g['compliance_pct']}%")
        c_pct.alignment = align_right
        c_pct.font = bold_font
        
        c_gap = ws1.cell(row=row_curr, column=12, value=f"{g['gap_pct']}%")
        c_gap.alignment = align_right
        
        c_stat = ws1.cell(row=row_curr, column=13, value=g["status_label"])
        c_stat.alignment = align_center
        
        if g["status_code"] == "OPTIMO":
            c_stat.fill = fill_green
            c_stat.font = font_green
        elif g["status_code"] == "EN_DESARROLLO":
            c_stat.fill = fill_yellow
            c_stat.font = font_yellow
        else:
            c_stat.fill = fill_red
            c_stat.font = font_red
            
        for c in range(1, 14):
            ws1.cell(row=row_curr, column=c).border = thin_border
            if row_curr % 2 == 0 and c != 13:
                ws1.cell(row=row_curr, column=c).fill = fill_zebra
                
        row_curr += 1
        
    # Fila de Totales
    ws1.cell(row=row_curr, column=1, value="TOTAL").font = bold_font
    ws1.cell(row=row_curr, column=2, value="CONSOLIDADO INSTITUCIONAL").font = bold_font
    ws1.cell(row=row_curr, column=3, value="100.0%").font = bold_font
    ws1.cell(row=row_curr, column=4, value=evaluation["total_items"]).font = bold_font
    ws1.cell(row=row_curr, column=5, value=evaluation["total_completo"]).font = bold_font
    ws1.cell(row=row_curr, column=6, value=evaluation["total_parcial"]).font = bold_font
    ws1.cell(row=row_curr, column=7, value=evaluation["total_no_cumple"]).font = bold_font
    ws1.cell(row=row_curr, column=8, value=evaluation["total_na"]).font = bold_font
    ws1.cell(row=row_curr, column=9, value=evaluation["total_points"]).font = bold_font
    ws1.cell(row=row_curr, column=10, value=evaluation["total_max_points"]).font = bold_font
    ws1.cell(row=row_curr, column=11, value=f"{evaluation['global_compliance']}%").font = bold_font
    ws1.cell(row=row_curr, column=12, value=f"{evaluation['global_gap']}%").font = bold_font
    ws1.cell(row=row_curr, column=13, value=evaluation["verdict_title"].split()[1]).font = bold_font
    
    for c in range(1, 14):
        ws1.cell(row=row_curr, column=c).border = thin_border
        ws1.cell(row=row_curr, column=c).fill = fill_gray

    # ====================================================
    # HOJA 2: MATRIZ DETALLADA GAP
    # ====================================================
    ws2 = wb.create_sheet(title="Matriz Detallada GAP")
    ws2.views.sheetView[0].showGridLines = True
    
    ws2.merge_cells("A1:N1")
    ws2["A1"] = "MATRIZ MAESTRA DE EVALUACIÓN Y ANÁLISIS DE BRECHAS (GAP) - NEWCARE OPTIMUS HEALTH"
    ws2["A1"].font = Font(name=font_family, size=13, bold=True, color="FFFFFF")
    ws2["A1"].fill = fill_navy
    ws2["A1"].alignment = align_center
    
    headers_m = [
        "ID", "Grupo / Capítulo", "Num.", "Título del Requisito", "Requisito Normativo Verificable",
        "Qué Audita New Health", "Fase Auditoría", "Puntos (0-2-NA)", "Calificación",
        "Evidencia Documental / Soporte", "Hallazgos / Brechas", "Plan de Acción Inmediato",
        "Responsable", "Prioridad", "Estado Tarea"
    ]
    
    for col_idx, h in enumerate(headers_m, 1):
        cell = ws2.cell(row=2, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border
        
    row_m = 3
    for it in items:
        ws2.cell(row=row_m, column=1, value=it["id"]).alignment = align_center
        ws2.cell(row=row_m, column=2, value=it.get("group_id", "")).alignment = align_center
        ws2.cell(row=row_m, column=3, value=it.get("code", "")).alignment = align_center
        ws2.cell(row=row_m, column=4, value=it.get("title", "")).alignment = align_left
        ws2.cell(row=row_m, column=5, value=it.get("requirement", "")).alignment = align_left
        ws2.cell(row=row_m, column=6, value=it.get("audit_criterion", "")).alignment = align_left
        ws2.cell(row=row_m, column=7, value=it.get("phase", "Certificación")).alignment = align_center
        
        sc = it.get("score", 0)
        c_score = ws2.cell(row=row_m, column=8, value=sc)
        c_score.alignment = align_center
        c_score.font = bold_font
        
        c_status = ws2.cell(row=row_m, column=9)
        c_status.alignment = align_center
        if sc == 2 or sc == "2":
            c_status.value = "Cumple Completo"
            c_status.fill = fill_green
            c_status.font = font_green
        elif sc == 1 or sc == "1":
            c_status.value = "Cumple Parcial"
            c_status.fill = fill_yellow
            c_status.font = font_yellow
        elif sc == "NA" or sc == "N":
            c_status.value = "No Aplica"
            c_status.fill = fill_gray
            c_status.font = font_gray
        else:
            c_status.value = "No Cumple"
            c_status.fill = fill_red
            c_status.font = font_red
            
        ws2.cell(row=row_m, column=10, value=it.get("evidence", "")).alignment = align_left
        ws2.cell(row=row_m, column=11, value=it.get("findings", "")).alignment = align_left
        ws2.cell(row=row_m, column=12, value=it.get("action", "")).alignment = align_left
        ws2.cell(row=row_m, column=13, value=it.get("responsible", "")).alignment = align_left
        ws2.cell(row=row_m, column=14, value=it.get("priority", "Media")).alignment = align_center
        ws2.cell(row=row_m, column=15, value=it.get("status", "Pendiente")).alignment = align_center
        
        for c in range(1, 16):
            ws2.cell(row=row_m, column=c).border = thin_border
            if row_m % 2 == 0 and c != 9:
                ws2.cell(row=row_m, column=c).fill = fill_zebra
                
        row_m += 1

    # ====================================================
    # HOJA 3: PLAN DE MEJORA CONTINUA (ACTION TRACKER)
    # ====================================================
    ws3 = wb.create_sheet(title="Plan de Mejora Continua")
    ws3.views.sheetView[0].showGridLines = True
    
    ws3.merge_cells("A1:H1")
    ws3["A1"] = "PLAN DE ACCIÓN Y MEJORA CONTINUA PARA EL CIERRE DE BRECHAS (AUDITORÍA NEW HEALTH)"
    ws3["A1"].font = Font(name=font_family, size=13, bold=True, color="FFFFFF")
    ws3["A1"].fill = fill_teal
    ws3["A1"].alignment = align_center
    
    headers_p = [
        "Cód. Req.", "Capítulo", "Requisito con Brecha", "Hallazgo / Causa Raíz",
        "Acción Correctiva / Mejora", "Responsable Institucional", "Prioridad", "Estado"
    ]
    
    for col_idx, h in enumerate(headers_p, 1):
        cell = ws3.cell(row=2, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border
        
    row_p = 3
    for it in items:
        sc = it.get("score", 0)
        # Solo incluir ítems con brecha (0 o 1)
        if sc in [0, 1, "0", "1"]:
            ws3.cell(row=row_p, column=1, value=it.get("code", "")).alignment = align_center
            ws3.cell(row=row_p, column=2, value=it.get("group_id", "")).alignment = align_center
            ws3.cell(row=row_p, column=3, value=it.get("title", "")).alignment = align_left
            ws3.cell(row=row_p, column=4, value=it.get("findings", it.get("default_findings", ""))).alignment = align_left
            ws3.cell(row=row_p, column=5, value=it.get("action", it.get("default_action", ""))).alignment = align_left
            ws3.cell(row=row_p, column=6, value=it.get("responsible", it.get("default_responsible", ""))).alignment = align_left
            
            prio = it.get("priority", "Media")
            c_prio = ws3.cell(row=row_p, column=7, value=prio)
            c_prio.alignment = align_center
            if prio == "Alta":
                c_prio.font = Font(name=font_family, size=10, bold=True, color="DC2626")
            elif prio == "Media":
                c_prio.font = Font(name=font_family, size=10, bold=True, color="D97706")
                
            ws3.cell(row=row_p, column=8, value=it.get("status", "Pendiente")).alignment = align_center
            
            for c in range(1, 9):
                ws3.cell(row=row_p, column=c).border = thin_border
                if row_p % 2 == 0:
                    ws3.cell(row=row_p, column=c).fill = fill_zebra
            row_p += 1

    # Ajuste automático del ancho de columnas
    for sheet in [ws1, ws2, ws3]:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                lines = val.splitlines()
                if lines:
                    val = max(lines, key=len)
                if len(val) > max_len:
                    max_len = len(val)
            # Limitar anchos para columnas de texto largo
            adjusted_width = min(max(max_len + 3, 10), 50)
            sheet.column_dimensions[col_letter].width = adjusted_width

    # Guardar en buffer binario
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()