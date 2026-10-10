import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_user_manual():
    doc = Document()

    # Configuración de Márgenes
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Paleta de Colores Institucionales Optimus Health & NewCare
    NAVY = RGBColor(7, 55, 99)       # 073763
    SAPPHIRE = RGBColor(28, 69, 135) # 1C4587
    STEEL = RGBColor(46, 117, 181)   # 2E75B5
    DARK = RGBColor(30, 41, 59)      # 1E293B
    MUTED = RGBColor(100, 116, 139)  # 64748B

    def set_cell_shading(cell, color_hex):
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shd)

    def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    # -------------------------------------------------------------------------
    # ENCABEZADO INSTITUCIONAL CON LOGO
    # -------------------------------------------------------------------------
    tbl_hdr = doc.add_table(rows=1, cols=2)
    tbl_hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_hdr.autofit = False
    
    # Celda izquierda: Logo
    c_logo = tbl_hdr.cell(0, 0)
    c_logo.width = Inches(2.2)
    p_logo = c_logo.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if os.path.exists("Logo Optimus Health.jpeg"):
        try:
            p_logo.add_run().add_picture("Logo Optimus Health.jpeg", width=Inches(1.8))
        except Exception:
            p_logo.add_run("OPTIMUS HEALTH IPS").bold = True

    # Celda derecha: Títulos
    c_tit = tbl_hdr.cell(0, 1)
    c_tit.width = Inches(4.6)
    p_tit = c_tit.paragraphs[0]
    p_tit.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    r_sub = p_tit.add_run("MODELO DE ATENCIÓN NEWCARE • FUNDACIÓN NEW HEALTH\n")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(8.5)
    r_sub.font.bold = True
    r_sub.font.color.rgb = SAPPHIRE

    r_main = p_tit.add_run("MANUAL DE USUARIO DEL EVALUADOR\n")
    r_main.font.name = "Calibri"
    r_main.font.size = Pt(16)
    r_main.font.bold = True
    r_main.font.color.rgb = NAVY

    r_doc = p_tit.add_run("Herramienta Automatizada de Selección, Baremación y Auditoría de Idoneidad del Talento Humano")
    r_doc.font.name = "Calibri"
    r_doc.font.size = Pt(9.5)
    r_doc.font.italic = True
    r_doc.font.color.rgb = MUTED

    doc.add_paragraph() # Espacio

    # -------------------------------------------------------------------------
    # FICHA TÉCNICA DEL MANUAL
    # -------------------------------------------------------------------------
    tbl_meta = doc.add_table(rows=2, cols=4)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_headers = ["Documento:", "Versión:", "Destinatarios:", "Sede Institucional:"]
    meta_values = [
        "MN-TH-NC-001 (Manual de Usuario)",
        "3.0 (Automatización Total)",
        "Comité de Selección, Talento Humano, Dirección Médica y Evaluadores",
        "Optimus Health IPS • Sincelejo, Sucre (Colombia)"
    ]

    for col_idx in range(4):
        c_h = tbl_meta.cell(0, col_idx)
        set_cell_shading(c_h, "073763")
        set_cell_margins(c_h, top=80, bottom=80, left=100, right=100)
        p = c_h.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(meta_headers[col_idx])
        r.font.name = "Calibri"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

        c_v = tbl_meta.cell(1, col_idx)
        set_cell_shading(c_v, "F1F5F9")
        set_cell_margins(c_v, top=80, bottom=80, left=100, right=100)
        p_v = c_v.paragraphs[0]
        p_v.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_v = p_v.add_run(meta_values[col_idx])
        r_v.font.name = "Calibri"
        r_v.font.size = Pt(8.5)
        r_v.font.color.rgb = DARK

    doc.add_paragraph()

    # Helper para Títulos
    def add_custom_heading(text, level=1):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(text)
        r.font.name = "Calibri"
        r.font.bold = True
        if level == 1:
            r.font.size = Pt(13)
            r.font.color.rgb = NAVY
        elif level == 2:
            r.font.size = Pt(11)
            r.font.color.rgb = SAPPHIRE
        elif level == 3:
            r.font.size = Pt(10)
            r.font.color.rgb = STEEL
        return h

    def add_callout(text, alert_type="INFO"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = tbl.cell(0, 0)
        c.width = Inches(6.8)
        set_cell_margins(c, top=100, bottom=100, left=150, right=150)
        
        bg_col = "E8EEF5" if alert_type == "INFO" else ("FFF2CC" if alert_type == "WARNING" else "E2F0D9")
        border_col = "1C4587" if alert_type == "INFO" else ("B25900" if alert_type == "WARNING" else "276A3C")
        icon_txt = "ℹ️ INFORMACIÓN IMPORTANTE:" if alert_type == "INFO" else ("⚠️ REGLA DE ORO / ATENCIÓN:" if alert_type == "WARNING" else "✅ BUENA PRÁCTICA RECOMENDADA:")
        
        set_cell_shading(c, bg_col)
        p = c.paragraphs[0]
        r_tit = p.add_run(f"{icon_txt}\n")
        r_tit.font.name = "Calibri"
        r_tit.font.size = Pt(9.5)
        r_tit.font.bold = True
        r_tit.font.color.rgb = NAVY if alert_type == "INFO" else (RGBColor(178, 89, 0) if alert_type == "WARNING" else RGBColor(39, 106, 60))

        r_txt = p.add_run(text)
        r_txt.font.name = "Calibri"
        r_txt.font.size = Pt(9)
        r_txt.font.color.rgb = DARK
        doc.add_paragraph()

    # -------------------------------------------------------------------------
    # SECCIÓN 1: INTRODUCCIÓN Y PROPÓSITO
    # -------------------------------------------------------------------------
    add_custom_heading("1. ¿Para qué sirve esta herramienta?", 1)
    doc.add_paragraph(
        "El archivo de Excel 'Baremacion_Personal_Asistencial_NEWCARE_Optimus_Health.xlsx' es el instrumento tecnológico "
        "diseñado para evaluar, calificar y auditar de forma 100% objetiva a los profesionales y técnicos que aspiran o "
        "hacen parte del equipo tratante del Programa NewCare en Optimus Health IPS (Sincelejo, Sucre)."
    )
    doc.add_paragraph(
        "Este modelo responde a los estándares del Método NEWPALEX® de la Fundación New Health (Criterio 2.4.2 y Cap. 3) "
        "y al Sistema Único de Habilitación en Salud de Colombia (Resolución 3100 de 2019), garantizando que todo el equipo "
        "cuente con la idoneidad clínica, la formación en cuidados paliativos y cronicidad, y la calidez humana requeridas."
    )

    add_callout(
        "El sistema calcula las notas de forma 100% automática mediante fórmulas interconectadas. El evaluador NO debe "
        "calcular promedios a mano ni sobreescribir celdas con fórmulas. Solo debe ingresar datos en las casillas correspondientes.",
        "INFO"
    )

    # -------------------------------------------------------------------------
    # SECCIÓN 2: MAPA DE LAS 6 HOJAS DE EXCEL
    # -------------------------------------------------------------------------
    add_custom_heading("2. Mapa de Navegación del Libro de Excel (¿Qué hay en cada hoja?)", 1)
    doc.add_paragraph(
        "El libro de trabajo está organizado en 6 hojas de cálculo con colores institucionales y funciones claramente delimitadas:"
    )

    hojas_desc = [
        ("1. Metodologia NewCare", "Marco conceptual e institucional. Explica la arquitectura triádica de evaluación (100 puntos) y la escala de dictamen oficial (Sobresaliente, Condicionado o No Idóneo). Es una hoja puramente informativa de consulta."),
        ("2. Filtros Excluyentes SI-NO", "Filtro legal y habilitante obligatorio (Gatekeeper). Contiene los 9 requisitos mínimos normativos (ReTHUS, Tarjeta Profesional, Violencia Sexual, BLS/ACLS, etc.). Si un candidato incumple uno solo, queda descartado."),
        ("3. Baremo CV (50 Puntos)", "Calificador y auditor de Hoja de Vida. Evalúa 8 dimensiones y 12 subcriterios aplicando estrictamente el Tope Rubro. Cuenta con un visor superior individual y una matriz inferior con los méritos de cada profesional."),
        ("4. Psicotecnico y Entrevista", "Evaluación psico-relacional (50 puntos máximos). Califica pruebas de Personalidad Resistente (CPR sobre 12 pts), Actitudes Funcionales (DAS sobre 8 pts) y la Entrevista por 9 Competencias NewCare (sobre 30 pts)."),
        ("5. Fichas de Cargo NewCare", "Catálogo de perfiles institucionales según el Organigrama de Optimus Health. Describe las funciones clínicas, requisitos y complejidad asignada para Médicos, Enfermeros, Terapeutas, Psicólogos, etc."),
        ("6. Registro y Evaluador", "Consolidado Ejecutivo y Cuadro de Mando en Vivo. Centraliza la nota final de todos los candidatos (sobre 100), emite el dictamen institucional y actualiza los indicadores (KPIs) de aprobación en tiempo real.")
    ]

    tbl_hojas = doc.add_table(rows=len(hojas_desc)+1, cols=2)
    tbl_hojas.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_hojas.cell(0, 0).width = Inches(2.2)
    tbl_hojas.cell(0, 1).width = Inches(4.6)
    
    set_cell_shading(tbl_hojas.cell(0, 0), "073763")
    set_cell_shading(tbl_hojas.cell(0, 1), "073763")
    r_h1 = tbl_hojas.cell(0, 0).paragraphs[0].add_run("Nombre de la Hoja")
    r_h1.font.bold = True
    r_h1.font.color.rgb = RGBColor(255, 255, 255)
    r_h2 = tbl_hojas.cell(0, 1).paragraphs[0].add_run("Función y Contenido Principal")
    r_h2.font.bold = True
    r_h2.font.color.rgb = RGBColor(255, 255, 255)

    for i, (nombre_h, desc_h) in enumerate(hojas_desc, start=1):
        c1 = tbl_hojas.cell(i, 0)
        c2 = tbl_hojas.cell(i, 1)
        c1.width = Inches(2.2)
        c2.width = Inches(4.6)
        set_cell_margins(c1, top=60, bottom=60, left=80, right=80)
        set_cell_margins(c2, top=60, bottom=60, left=80, right=80)
        if i % 2 == 0:
            set_cell_shading(c1, "F8FAFC")
            set_cell_shading(c2, "F8FAFC")
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(nombre_h)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9)
        r1.font.bold = True
        r1.font.color.rgb = NAVY

        p2 = c2.paragraphs[0]
        r2 = p2.add_run(desc_h)
        r2.font.name = "Calibri"
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = DARK

    doc.add_paragraph()

    # -------------------------------------------------------------------------
    # SECCIÓN 3: GUÍA OPERATIVA PASO A PASO PARA EVALUAR UN CANDIDATO
    # -------------------------------------------------------------------------
    add_custom_heading("3. Flujo Operativo: ¿Cómo calificar a un candidato paso a paso?", 1)
    doc.add_paragraph(
        "El proceso de calificación sigue un orden lógico secuencial. Para registrar y calificar a un profesional, "
        "siga estrictamente estos 5 pasos:"
    )

    pasos = [
        ("PASO 1: Registrar el Candidato en la Hoja 6", 
         "Abra la hoja '6. Registro y Evaluador'. Diríjase a una fila disponible (por ejemplo, la fila 5 si ya está, o la fila 21 si es un candidato nuevo). Escriba su Nombre completo en la Columna B y el Cargo Postulado en la Columna C. Al ingresar el nombre, el profesional se incorporará automáticamente a las listas desplegables del sistema."),
        
        ("PASO 2: Verificar Filtros Habilitantes en la Hoja 2", 
         "Vaya a la hoja '2. Filtros Excluyentes SI-NO'. En la celda C4, use el menú desplegable para elegir al candidato. Revise uno por uno los 9 requisitos normativos (diploma, ReTHUS, violencia sexual, BLS, etc.). En la Columna G debe indicar si 'CUMPLE' o 'NO CUMPLE'. Si el candidato incumple un solo filtro, el sistema emitirá en G17 'NO CUMPLE (DESCARTADO)' y en la Hoja 6 su puntaje final se anulará a 0 automáticamente."),

        ("PASO 3: Auditar la Hoja de Vida en la Hoja 3 (Entender el 'Tope Rubro')", 
         "Diríjase a la hoja '3. Baremo CV (50 Puntos)'. En la celda C3 elija al candidato. En la tabla inferior (filas 22 a 48), ubique la fila del candidato e ingrese las cantidades reales acreditadas: años generales, años en paliativos, especialidades, cursos, publicaciones, etc.\n"
         "Al volver a la tabla superior (filas 5 a 16), verá cómo opera el 'Tope Rubro': la Columna G muestra lo acreditado, la Columna H calcula los puntos brutos, y la Columna I aplica el corte máximo permitido. La Columna J le confirmará si el candidato superó el tope con una alerta amarilla que indica los puntos excedentes no computados."),

        ("PASO 4: Ingresar Psicotécnico y Entrevista en la Hoja 4", 
         "Abra la hoja '4. Psicotecnico y Entrevista'. En la matriz inferior (filas 25 a 51), ubique al candidato:\n"
         "• Columna D: Ingrese el puntaje de Personalidad Resistente CPR (Máx 12.0).\n"
         "• Columna E: Ingrese el puntaje de Actitudes Funcionales DAS (Máx 8.0).\n"
         "• Columna H: Ingrese la nota promedio de la Entrevista por Competencias (escala 0.0 a 10.0).\n"
         "El sistema calculará automáticamente el subtotal psicotécnico (Col F), el subtotal de entrevista ponderada (Col I) y el total psico-relacional (Col J). Puede consultar su ficha en vivo en las celdas K5:M10."),

        ("PASO 5: Consultar Dictamen Final y Cuadro de Mando en la Hoja 6", 
         "Regrese a la hoja '6. Registro y Evaluador'. Verá que todas las columnas (D, E, F, G, H, I) se han actualizado sin que usted tenga que digitar nada:\n"
         "• Col E: Muestra la nota del CV traída de la Hoja 3.\n"
         "• Col F: Muestra la nota del Psicotécnico traída de la Hoja 4.\n"
         "• Col G: Muestra la nota de la Entrevista traída de la Hoja 4.\n"
         "• Col H: Suma el Puntaje Final sobre 100 puntos.\n"
         "• Col I: Emite el dictamen oficial institucional con color automático.")
    ]

    for titulo_paso, detalle_paso in pasos:
        add_custom_heading(titulo_paso, 2)
        doc.add_paragraph(detalle_paso)

    add_callout(
        "REGLA CRÍTICA DE INTEGRIDAD: NUNCA digite valores numéricos en las columnas E, F, G, H e I de la Hoja 6. "
        "Esas celdas contienen fórmulas dinámicas (VLOOKUP / BUSCARV). Si escribe un número encima, romperá el vínculo automático "
        "con las hojas de evaluación.",
        "WARNING"
    )

    # -------------------------------------------------------------------------
    # SECCIÓN 4: COMPRENSIÓN CLAVE DEL "TOPE RUBRO"
    # -------------------------------------------------------------------------
    add_custom_heading("4. ¿Qué es el 'Tope Rubro' y cómo interpretarlo en la auditoría?", 1)
    doc.add_paragraph(
        "El concepto de 'Tope Rubro' es el corazón metodológico exigido por la Fundación New Health. Comprenderlo evita "
        "errores en las auditorías de acreditación:"
    )
    doc.add_paragraph(
        "• Propósito: Evita que un candidato compense su falta de experiencia específica en cuidados paliativos acumulando "
        "muchos años de experiencia hospitalaria general. Cada subcriterio tiene un 'techo' que no se puede rebasar."
    )
    doc.add_paragraph(
        "• Ejemplo Práctico (Fila 5 del Baremo): El subcriterio de experiencia general otorga 1.0 punto por año, pero tiene un "
        "Tope Rubro de 7.0 puntos. Si el Dr. Julio Herrera acredita 10 años, el cálculo bruto da 10.0 puntos (Columna H), pero "
        "el puntaje asignado se corta estrictamente en 7.0 puntos (Columna I). La Columna J muestra: 'TOPE APLICADO (Exceso: 3.0 pts no contabilizados)'."
    )
    doc.add_paragraph(
        "• En la Hoja de Vida, el aspirante solo alcanzará los 15.0 puntos de Experiencia si acredita además experiencia específica "
        "en cronicidad/paliativos (hasta 8.0 puntos adicionales a 2.0 pts por año)."
    )

    # -------------------------------------------------------------------------
    # SECCIÓN 5: ESCALA INSTITUCIONAL DE DICTAMEN DE IDONEIDAD
    # -------------------------------------------------------------------------
    add_custom_heading("5. Escala de Decisión y Dictamen Institucional de Selección", 1)
    doc.add_paragraph(
        "Al finalizar la evaluación de los 100 puntos, la Columna I de la Hoja 6 emite automáticamente uno de los siguientes dictámenes:"
    )

    dictamenes_info = [
        ("IDÓNEO SOBRESALIENTE", ">= 85.0 Puntos", "Verde", "El candidato cumple con el estándar de excelencia NewCare. Asignación inmediata al equipo tratante sin condicionamientos."),
        ("IDÓNEO CONDICIONADO", "70.0 a 84.9 Puntos", "Amarillo", "Apto para contratación. Requiere plan intensivo de formación y entrenamiento prioritario NewCare durante los primeros 90 días."),
        ("NO IDÓNEO", "< 70.0 Puntos", "Rojo", "No alcanza el umbral de idoneidad exigido por los estándares New Health. Queda descartado del proceso."),
        ("DESCARTADO EN FILTROS", "Filtro = NO CUMPLE", "Rojo", "Incumple uno o más requisitos legales mínimos habilitantes de la Hoja 2. Su puntaje se anula a 0 automáticamente.")
    ]

    tbl_dict = doc.add_table(rows=len(dictamenes_info)+1, cols=4)
    tbl_dict.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_shading(tbl_dict.cell(0, 0), "073763")
    set_cell_shading(tbl_dict.cell(0, 1), "073763")
    set_cell_shading(tbl_dict.cell(0, 2), "073763")
    set_cell_shading(tbl_dict.cell(0, 3), "073763")

    for col_i, tit in enumerate(["Dictamen Oficial", "Rango de Puntaje", "Alerta Visual", "Acción Institucional"]):
        r = tbl_dict.cell(0, col_i).paragraphs[0].add_run(tit)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for idx, (nom_d, rng_d, col_d, acc_d) in enumerate(dictamenes_info, start=1):
        c0 = tbl_dict.cell(idx, 0)
        c1 = tbl_dict.cell(idx, 1)
        c2 = tbl_dict.cell(idx, 2)
        c3 = tbl_dict.cell(idx, 3)
        set_cell_margins(c0, top=60, bottom=60, left=80, right=80)
        set_cell_margins(c1, top=60, bottom=60, left=80, right=80)
        set_cell_margins(c2, top=60, bottom=60, left=80, right=80)
        set_cell_margins(c3, top=60, bottom=60, left=80, right=80)
        
        c0.paragraphs[0].add_run(nom_d).bold = True
        c1.paragraphs[0].add_run(rng_d)
        c2.paragraphs[0].add_run(col_d)
        c3.paragraphs[0].add_run(acc_d)
        
        bg_row = "E2F0D9" if "SOBRESALIENTE" in nom_d else ("FFF2CC" if "CONDICIONADO" in nom_d else "FCE4D6")
        set_cell_shading(c0, bg_row)
        set_cell_shading(c1, bg_row)
        set_cell_shading(c2, bg_row)
        set_cell_shading(c3, bg_row)

    doc.add_paragraph()

    # -------------------------------------------------------------------------
    # SECCIÓN 6: PREGUNTAS FRECUENTES Y RESOLUCIÓN DE DUDAS (FAQ)
    # -------------------------------------------------------------------------
    add_custom_heading("6. Preguntas Frecuentes y Solución de Problemas (FAQ)", 1)

    faqs = [
        ("¿Cómo agrego un nuevo candidato para que aparezca en los menús desplegables?",
         "1. Vaya a la hoja '6. Registro y Evaluador' y escriba el nombre en una fila vacía de la Columna B (ej. B21).\n"
         "2. En Excel, vaya a la pestaña superior 'Fórmulas' > 'Administrador de nombres'.\n"
         "3. Seleccione el nombre 'Lista_Candidatos' y en 'Se refiere a:' extienda el rango hasta la fila correspondiente (por ejemplo: ='6. Registro y Evaluador'!$B$5:$B$25).\n"
         "4. Al guardar, el nuevo candidato aparecerá en todos los menús desplegables del libro."),

        ("¿Por qué en la Hoja 6 un candidato tiene puntaje 0 si obtuvo buenas notas en CV y entrevista?",
         "Revise de inmediato la Columna D de la Hoja 6. Si dice 'NO CUMPLE (DESCARTADO)', significa que en la Hoja 2 (Filtros Excluyentes) se marcó al menos un requisito como 'NO CUMPLE'. Por norma legal de habilitación, un candidato con filtros incompletos no puede obtener puntaje."),

        ("¿Qué hago si una celda muestra '#N/A' o '#¡VALOR!'?",
         "Verifique que el nombre del candidato en la celda del menú desplegable esté escrito exactamente igual al que figura en la matriz inferior. Al usar la lista desplegable se evitan errores tipográficos o espacios adicionales."),

        ("¿Cómo imprimo o presento los resultados para una auditoría de la Fundación New Health?",
         "Para presentar las evidencias de selección ante los auditores internacionales:\n"
         "• La Hoja 6 consolidada sirve como Acta Ejecutiva de Calificación de todo el equipo asistencial.\n"
         "• La Hoja 3 (con el candidato seleccionado en C3) sirve como Expediente Individual de Baremación Curricular.\n"
         "• La Hoja 2 sirve como Ficha de Habilitación Normativa Obligatoria.")
    ]

    for q_tit, q_ans in faqs:
        add_custom_heading(q_tit, 2)
        doc.add_paragraph(q_ans)

    add_callout(
        "CONSEJO DE SEGURIDAD: Guarde periódicamente copias de seguridad del archivo con fecha (por ejemplo: "
        "'Baremacion_Personal_Asistencial_2026_10_08.xlsx') antes de iniciar una nueva convocatoria de selección.",
        "INFO"
    )

    out_file = "MANUAL_DE_USUARIO_SISTEMA_BAREMACION_NEWCARE_OPTIMUS_HEALTH.docx"
    doc.save(out_file)
    print(f"Manual de usuario generado exitosamente en: {out_file}")

if __name__ == "__main__":
    create_user_manual()
