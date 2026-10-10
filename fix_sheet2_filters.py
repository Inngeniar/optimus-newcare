import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill, Font

def fix_sheet2_filters():
    wb_path = "Baremacion_Personal_Asistencial_NEWCARE_Optimus_Health.xlsx"
    wb = openpyxl.load_workbook(wb_path)

    ws2 = wb["2. Filtros Excluyentes SI-NO"]

    # 1. Asegurar validación de lista de candidatos en C4
    # Verificar si ya existe, si no, crearla
    dv_cand = None
    for dv in ws2.data_validations.dataValidation:
        if "C4" in str(dv.sqref):
            dv_cand = dv
            break

    if not dv_cand:
        dv_cand = DataValidation(type="list", formula1="=Lista_Candidatos", allow_blank=False)
        ws2.add_data_validation(dv_cand)
        dv_cand.add("C4")

    # 2. Agregar validación desplegable para G7:G15 (CUMPLE / NO CUMPLE)
    dv_flt = DataValidation(
        type="list",
        formula1='"CUMPLE,NO CUMPLE"',
        allow_blank=False
    )
    dv_flt.promptTitle = "Verificación Habilitante"
    dv_flt.prompt = "Seleccione si el candidato CUMPLE o NO CUMPLE con el requisito."
    dv_flt.errorTitle = "Opción No Válida"
    dv_flt.error = "Debe seleccionar 'CUMPLE' o 'NO CUMPLE'."
    dv_flt.showErrorMessage = True
    dv_flt.showInputMessage = True

    ws2.add_data_validation(dv_flt)
    dv_flt.add("G7:G15")
    print("Validación de datos 'CUMPLE,NO CUMPLE' agregada exitosamente a G7:G15.")

    # 3. Formato Condicional en G7:G15 y G17
    green_fill = PatternFill(start_color='E2F0D9', end_color='E2F0D9', fill_type='solid')
    green_font = Font(name='Calibri', size=10, bold=True, color='276A3C')

    red_fill = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')
    red_font = Font(name='Calibri', size=10, bold=True, color='C00000')

    ws2.conditional_formatting.add('G7:G15', FormulaRule(formula=['EXACT(G7,"CUMPLE")'], fill=green_fill, font=green_font))
    ws2.conditional_formatting.add('G7:G15', FormulaRule(formula=['EXACT(G7,"NO CUMPLE")'], fill=red_fill, font=red_font))

    ws2.conditional_formatting.add('G17', FormulaRule(formula=['ISNUMBER(SEARCH("100%",G17))'], fill=green_fill, font=green_font))
    ws2.conditional_formatting.add('G17', FormulaRule(formula=['ISNUMBER(SEARCH("NO CUMPLE",G17))'], fill=red_fill, font=red_font))

    wb.save(wb_path)
    print("Hoja 2 guardada exitosamente con validaciones y formatos condicionales activos.")

if __name__ == "__main__":
    fix_sheet2_filters()
