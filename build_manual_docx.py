import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = Document()

# Margenes estandar
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

TEAL = RGBColor(15, 118, 110)
SLATE = RGBColor(51, 65, 85)

def set_cell_shading(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def add_header_box(title, subtitle):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_shading(cell, "0F766E")
    cell.width = Inches(6.8)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_t = p.add_run("OPTIMUS HEALTH IPS  •  PROGRAMA NEWCARE\n")
    run_t.font.name = "Segoe UI"
    run_t.font.size = Pt(11)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(255, 255, 255)
    
    run_main = p.add_run(f"{title}\n")
    run_main.font.name = "Segoe UI"
    run_main.font.size = Pt(18)
    run_main.font.bold = True
    run_main.font.color.rgb = RGBColor(255, 255, 255)
    
    run_sub = p.add_run(subtitle)
    run_sub.font.name = "Segoe UI"
    run_sub.font.size = Pt(10)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(204, 251, 241)
    doc.add_paragraph()

add_header_box("MANUAL DE USUARIO Y GUÍA DE DESPLIEGUE", "Herramienta de Análisis GAP para la Certificación NewCare (Fundación New Health)")

p_meta = doc.add_paragraph()
r1 = p_meta.add_run("Versión del Sistema: ")
r1.bold = True
p_meta.add_run("2.0 (Edición Corporativa)\n")
r2 = p_meta.add_run("Destinatario: ")
r2.bold = True
p_meta.add_run("Patrocinador del Proyecto, Gerencia General y Equipo Asistencial Optimus Health IPS\n")
r3 = p_meta.add_run("Sede Asistencial: ")
r3.bold = True
p_meta.add_run("Sincelejo, Sucre • Colombia\n")
r4 = p_meta.add_run("Fecha de Emisión: ")
r4.bold = True
p_meta.add_run("Septiembre de 2026\n")

def add_heading_styled(text, level=1):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(4)
    run = h.add_run(text)
    run.font.name = "Segoe UI"
    run.font.bold = True
    if level == 1:
        run.font.size = Pt(14)
        run.font.color.rgb = TEAL
    elif level == 2:
        run.font.size = Pt(12)
        run.font.color.rgb = SLATE
    return h

add_heading_styled("1. Introducción y Propósito de la Aplicación", 1)
doc.add_paragraph(
    "La aplicación de Análisis GAP NewCare es una plataforma web desarrollada a la medida para Optimus Health IPS, "
    "con el objetivo de verificar, medir y optimizar el grado de cumplimiento frente a los estándares de certificación "
    "de la Fundación New Health para el programa de atención integral de Insuficiencias Orgánicas Crónicas Complejas "
    "(IOCC: Cardíaca, Respiratoria, Renal y Hepática avanzadas)."
)
doc.add_paragraph(
    "La herramienta adopta la metodología de Análisis GAP de ISO 9001:2015, evaluando 64 requisitos normativos distribuidos "
    "en los 13 capítulos del estándar, cuantificando la brecha (GAP) y generando automáticamente planes de mejora continua para "
    "garantizar la superación exitosa de las fases de Informe Favorable y Certificación Definitiva."
)

add_heading_styled("2. Guía de Instalación Rápida para el Patrocinador", 1)
doc.add_paragraph(
    "La solución está diseñada para ejecutarse en cualquier computador con sistema operativo Windows, garantizando que el "
    "patrocinador del proyecto o cualquier miembro directivo pueda operarla de forma autónoma sin configuraciones complejas."
)

add_heading_styled("2.1 Requisitos Previos", 2)
doc.add_paragraph("• Sistema Operativo: Windows 10 o Windows 11 (64 bits).")
doc.add_paragraph("• Python: Versión 3.10 o superior instalada (descargable gratis desde python.org). Importante: Marcar la casilla 'Add Python to PATH' durante su instalación.")
doc.add_paragraph("• Conexión a Internet: Requerida únicamente la primera vez para descargar las librerías.")

add_heading_styled("2.2 Instalación en 1 Clic", 2)
doc.add_paragraph("1. Descomprima o copie la carpeta del proyecto 'Proyecto NEWCARE' en el computador de destino.")
doc.add_paragraph("2. Haga doble clic en el archivo: INSTALADOR_AUTOMATICO.bat")
doc.add_paragraph("3. El instalador creará un entorno virtual aislado ('env_newcare') e instalará automáticamente todas las librerías necesarias (Streamlit, Plotly, Pandas, OpenPyXL, Python-Docx).")
doc.add_paragraph("4. Al finalizar la instalación, verá un mensaje en pantalla indicando: '¡INSTALACIÓN Y CONFIGURACIÓN COMPLETADA CON ÉXITO!'.")

add_heading_styled("2.3 Inicio de la Aplicación", 2)
doc.add_paragraph("Para abrir la aplicación en el día a día, simplemente:")
doc.add_paragraph("👉 Haga doble clic en el archivo: ejecutar_app.bat")
doc.add_paragraph("Se abrirá una ventana de comando que iniciará el servidor local y abrirá de inmediato su navegador web predeterminado en: http://localhost:8501.")

add_heading_styled("3. Módulos y Funcionalidades Principales", 1)

doc.add_paragraph("📊 1. Panel Ejecutivo (Dashboard):")
doc.add_paragraph(
    "Presenta las tarjetas de mando de alto nivel: Calificación Global Ponderada (con semáforo de viabilidad), Brecha Total (GAP), "
    "Meta de Certificación y Diagnóstico Global. Incluye el Gráfico Radial / Spider Chart de los 13 capítulos, el desglose de barras "
    "horizontales y la tabla ejecutiva de cumplimiento. Alerta sobre las brechas críticas que impiden la certificación."
)

doc.add_paragraph("📝 2. Estudio de Evaluación GAP (Requisito por Requisito):")
doc.add_paragraph(
    "Permite auditar individualmente los 64 requisitos normativos con botones de un clic (0: No Cumple, 1: Cumple Parcial, 2: Cumple Completo, NA: No Aplica). "
    "Cada requisito cuenta con un acordeón donde se registran las evidencias documentales, los hallazgos identificados, el plan de trabajo, el responsable y la prioridad. "
    "Incluye un botón interactivo para subir archivos de soporte (PDF, Word, Excel, fotos) que se almacenan directamente en la carpeta local de evidencias."
)

doc.add_paragraph("🎯 3. Plan de Mejora Continua:")
doc.add_paragraph(
    "Módulo interactivo para gestionar las tareas correctivas derivadas de los requisitos con brecha (0 y 1). Ofrece 3 modos de visualización: "
    "Fichas de Acción Detalladas (texto 100% visible, sin cortes), Vista Agrupada por Capítulos en Acordeón, y Tabla Resumen Multilínea. "
    "Permite cambiar el estado de la tarea (Pendiente, En Proceso, Implementado) y marcar como cumplido (2 pts) para recalcular el porcentaje global en vivo."
)

doc.add_paragraph("🔍 4. Simulador de Auditoría:")
doc.add_paragraph(
    "Compara el estado de Optimus Health frente a los dos hitos de la Fundación New Health: "
    "a) Hito 1: Requisitos de Informe Favorable (autorización de inicio operativo), y "
    "b) Hito 2: Requisitos de Certificación Definitiva (Sello NewCare)."
)

doc.add_paragraph("📁 5. Repositorio Documental Central y Carga Transversal:")
doc.add_paragraph(
    "Centro unificado para almacenar manuales consolidados (Manual NewCare v2.0), instructivos oficiales de New Health, documentos de habilitación "
    "y mallas de indicadores transversales. Permite subir múltiples documentos a la vez y vincularlos con un solo clic como evidencia a cualquier requisito."
)

doc.add_paragraph("📚 6. Manual y Referencias:")
doc.add_paragraph(
    "Biblioteca técnica con resúmenes de las Guías Clínicas adoptadas (AHA Falla Cardíaca, GOLD EPOC, KDIGO Renal, Cirrosis), "
    "escalas clínicas parametrizadas en INFORMEX (NECPAL, Barthel, Charlson, Zarit, Pfeiffer, Braden) y glosario terminológico oficial."
)

doc.add_paragraph("📥 7. Centro de Exportación Multiformato:")
doc.add_paragraph(
    "Permite descargar los resultados en 6 formatos corporativos: Excel Oficial (.xlsx con 3 hojas semaforizadas), Documento Word Oficial (.docx listo para firmas directivas), "
    "Plan de Mejora en CSV (UTF-8 con BOM), Matriz Maestra en CSV, Copia de Respaldo JSON y Reporte Ejecutivo HTML imprimible."
)

add_heading_styled("4. Gestión de Datos, Guardado y Respaldos", 1)
doc.add_paragraph(
    "• Guardado Automático Local: Todas las calificaciones, evidencias y notas se guardan de forma instantánea en el archivo: "
    "Proyecto NEWCARE/data/gap_assessment.json. La información nunca se pierde al cerrar el navegador o apagar el equipo."
)
doc.add_paragraph(
    "• Respaldos Manuales: Desde el Centro de Exportación se puede descargar una Copia de Seguridad JSON con la fecha del día para trazabilidad histórica."
)
doc.add_paragraph(
    "• Botones de Restauración: En la barra lateral izquierda se dispone de botones para restaurar la Línea Base Inicial (↺ Base), "
    "limpiar todo a cero para una auditoría en blanco (🧹 Limpiar Todo), o simular el escenario de 100% de cumplimiento (🏆 Simular 100%)."
)

add_heading_styled("5. Firmas y Responsables de Aprobación", 1)
doc.add_paragraph("El presente manual y la herramienta han sido estructurados para la operación y certificación de Optimus Health IPS.")

tbl_sig = doc.add_table(rows=1, cols=3)
tbl_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
for row in tbl_sig.rows:
    for cell in row.cells:
        cell.width = Inches(2.2)

p0 = tbl_sig.cell(0, 0).paragraphs[0]
p0.text = "___________________________\nLíder Programa NEWCARE\nOptimus Health IPS"
p0.alignment = WD_ALIGN_PARAGRAPH.CENTER

p1 = tbl_sig.cell(0, 1).paragraphs[0]
p1.text = "___________________________\nCoordinador de Calidad\nOptimus Health IPS"
p1.alignment = WD_ALIGN_PARAGRAPH.CENTER

p2 = tbl_sig.cell(0, 2).paragraphs[0]
p2.text = "___________________________\nGerencia General / Patrocinador\nOptimus Health IPS"
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER

target_file = "MANUAL_DE_USUARIO_GAP_NEWCARE_OPTIMUS_HEALTH.docx"
doc.save(target_file)
print(f"Manual en Word guardado exitosamente en: {target_file}")
