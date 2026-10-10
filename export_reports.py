"""
export_reports.py
Modulo para la exportacion de informes en multiples formatos:
- Word (.docx) formal para directivos y comites
- CSV (.csv) con codificacion UTF-8-SIG para analisis y Excel
- JSON (.json) estructurado para respaldo y auditoria
"""

import io
import json
import datetime
from typing import List, Dict, Any
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def generate_docx_report(evaluation: Dict[str, Any], items: List[Dict[str, Any]]) -> bytes:
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    p_pre = doc.add_paragraph()
    p_pre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_inst = p_pre.add_run("OPTIMUS HEALTH IPS  •  SEDE SINCELEJO (SUCRE)\n")
    run_inst.font.name = "Arial"
    run_inst.font.size = Pt(13)
    run_inst.font.bold = True
    run_inst.font.color.rgb = RGBColor(15, 118, 110)

    run_prog = p_pre.add_run("PROGRAMA NEWCARE — ATENCION A LA CRONICIDAD COMPLEJA (IOCC)\n")
    run_prog.font.name = "Arial"
    run_prog.font.size = Pt(11)
    run_prog.font.bold = True
    run_prog.font.color.rgb = RGBColor(30, 41, 59)

    run_sub = p_pre.add_run("Informe Ejecutivo de Diagnostico GAP y Plan de Mejora Continua\nCertificacion ante la Fundacion New Health (Espana / Latinoamerica)")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(9.5)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph()

    p_meta = doc.add_paragraph()
    p_meta.add_run("Fecha de Generacion: ").bold = True
    p_meta.add_run(f"{datetime.date.today().strftime('%d/%m/%Y')}   |   ")
    p_meta.add_run("Metodologia: ").bold = True
    p_meta.add_run("Analisis de Brechas (GAP) ISO 9001:2015 adaptado a NEWCARE v2.0")

    doc.add_heading("1. Diagnostico Global de Cumplimiento", level=2)
    kpi_table = doc.add_table(rows=2, cols=4)
    kpi_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    kpi_headers = ["Cumplimiento Ponderado", "Brecha Global (GAP)", "Meta Certificacion", "Veredicto de Auditoria"]
    verdict_clean = evaluation["verdict_title"]
    for icon in ["🏆 ", "📑 ", "⚠️ ", "🚨 "]:
        verdict_clean = verdict_clean.replace(icon, "")
    kpi_values = [f"{evaluation['global_compliance']}%", f"{evaluation['global_gap']}%", f"{evaluation['target_pct']}%", verdict_clean]

    for i, h in enumerate(kpi_headers):
        cell = kpi_table.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(9)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_background(cell, "1E293B")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    for i, v in enumerate(kpi_values):
        cell = kpi_table.cell(1, i)
        cell.text = v
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(11)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_background(cell, "F8FAFC")
        if i == 0:
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(15, 118, 110)
        elif i == 1:
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(220, 38, 38)

    doc.add_paragraph()
    p_ver = doc.add_paragraph()
    p_ver.add_run("Diagnostico del Ente Auditor: ").bold = True
    p_ver.add_run(evaluation["verdict_desc"])

    doc.add_heading("2. Nivel de Cumplimiento por Capitulos Normativos (13 Grupos)", level=2)
    group_table = doc.add_table(rows=1, cols=7)
    group_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    g_headers = ["Cod", "Capitulo Normativo", "Peso %", "Req.", "Puntos", "% Cumpl.", "Estado"]

    for i, h in enumerate(g_headers):
        cell = group_table.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(9)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_background(cell, "334155")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    for g in evaluation["groups"]:
        row = group_table.add_row()
        row.cells[0].text = g["code"]
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row.cells[1].text = g["name"]
        row.cells[2].text = f"{g['weight_differential']}%"
        row.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        row.cells[3].text = str(g["total_items"])
        row.cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row.cells[4].text = f"{g['points_obtained']:.1f}/{int(g['max_points'])}"
        row.cells[4].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        row.cells[5].text = f"{g['compliance_pct']}%"
        row.cells[5].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        row.cells[5].paragraphs[0].runs[0].font.bold = True
        row.cells[6].text = g["status_label"]
        row.cells[6].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if g["status_code"] == "OPTIMO":
            set_cell_background(row.cells[6], "D1FAE5")
        elif g["status_code"] == "EN_DESARROLLO":
            set_cell_background(row.cells[6], "FEF3C7")
        else:
            set_cell_background(row.cells[6], "FEE2E2")
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8.5)

    doc.add_paragraph()
    doc.add_heading("3. Plan de Mejora Continua y Cierre de Brechas", level=2)
    p_desc = doc.add_paragraph(
        "A continuacion se detallan todas las tareas de mejora formuladas para cerrar las brechas "
        "identificadas en requisitos con calificacion de No Cumple (0%) o Cumple Parcial (50%). "
        "Este plan constituye la hoja de ruta obligatoria para presentar en la auditoria de New Health."
    )
    p_desc.runs[0].font.size = Pt(9.5)

    gap_items = [it for it in items if it.get("score") in [0, 1, "0", "1"]]
    groups_dict = {g["group_id"]: g for g in evaluation["groups"]}
    items_by_g = {}
    for it in gap_items:
        gid = it.get("group_id", "G0")
        if gid not in items_by_g:
            items_by_g[gid] = []
        items_by_g[gid].append(it)

    for gid, it_list in items_by_g.items():
        g_info = groups_dict.get(gid, {"name": gid, "code": ""})
        doc.add_heading(f"• {g_info['name']} ({len(it_list)} acciones)", level=3)
        for it in it_list:
            sc = it.get("score", 0)
            status_text = "No Cumple (0 pts)" if sc in [0, "0"] else "Cumple Parcial (1 pt)"
            t_box = doc.add_table(rows=4, cols=2)
            t_box.alignment = WD_TABLE_ALIGNMENT.CENTER

            t_box.cell(0, 0).text = f"Requisito: {it['code']} - {it['title']}"
            t_box.cell(0, 0).paragraphs[0].runs[0].font.bold = True
            set_cell_background(t_box.cell(0, 0), "F1F5F9")

            t_box.cell(0, 1).text = f"Estado: {status_text} | Prioridad: {it.get('priority', 'Media')}"
            t_box.cell(0, 1).paragraphs[0].runs[0].font.bold = True
            t_box.cell(0, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_cell_background(t_box.cell(0, 1), "F1F5F9")

            t_box.cell(1, 0).text = "Hallazgo / Brecha Detectada:"
            t_box.cell(1, 0).paragraphs[0].runs[0].font.bold = True
            t_box.cell(1, 1).text = it.get("findings", it.get("default_findings", "Pendiente de formalizar"))

            t_box.cell(2, 0).text = "Plan de Accion Inmediato:"
            t_box.cell(2, 0).paragraphs[0].runs[0].font.bold = True
            t_box.cell(2, 1).text = it.get("action", it.get("default_action", "Ejecutar adecuacion documental"))
            set_cell_background(t_box.cell(2, 1), "ECFDF5")

            t_box.cell(3, 0).text = f"Responsable: {it.get('responsible', it.get('default_responsible', 'Sin asignar'))}"
            t_box.cell(3, 0).paragraphs[0].runs[0].font.bold = True
            t_box.cell(3, 1).text = f"Estado de la Tarea: {it.get('status', 'Pendiente')}"
            t_box.cell(3, 1).paragraphs[0].runs[0].font.bold = True
            t_box.cell(3, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

            for row in t_box.rows:
                for c in row.cells:
                    for p in c.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(8.5)
            doc.add_paragraph()

    doc.add_paragraph()
    doc.add_heading("4. Firmas de Formalizacion y Compromiso Institucional", level=2)
    p_sig = doc.add_paragraph("El presente plan ha sido revisado y acordado por los responsables de area para garantizar el cumplimiento de los estandares de certificacion NewCare.")
    p_sig.runs[0].font.size = Pt(9.5)

    sig_table = doc.add_table(rows=2, cols=3)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell in sig_table.rows[0].cells:
        cell.text = "\n\n____________________________________"
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    sig_table.cell(1, 0).text = "Lider Programa NEWCARE\nOptimus Health IPS"
    sig_table.cell(1, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    sig_table.cell(1, 1).text = "Coordinacion de Calidad\nOptimus Health IPS"
    sig_table.cell(1, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    sig_table.cell(1, 2).text = "Direccion General / Gerencia\nOptimus Health IPS"
    sig_table.cell(1, 2).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    for row in sig_table.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()

def generate_csv_matrix(items: List[Dict[str, Any]]) -> bytes:
    import csv
    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow([
        "ID", "Grupo", "Numeral", "Titulo", "Requisito_Normativo",
        "Que_Audita_New_Health", "Fase", "Puntuacion", "Estado_Calificacion",
        "Evidencia_Documental", "Hallazgos_Brechas", "Plan_de_Accion",
        "Responsable", "Prioridad", "Estado_Tarea"
    ])
    for it in items:
        sc = it.get("score", 0)
        if sc in [2, "2"]:
            sc_text = "Cumple Completo"
        elif sc in [1, "1"]:
            sc_text = "Cumple Parcial"
        elif sc in ["NA", "N"]:
            sc_text = "No Aplica"
        else:
            sc_text = "No Cumple"
        writer.writerow([
            it.get("id", ""), it.get("group_id", ""), it.get("code", ""), it.get("title", ""),
            it.get("requirement", ""), it.get("audit_criterion", ""), it.get("phase", "Certificacion"),
            sc, sc_text, it.get("evidence", ""), it.get("findings", ""), it.get("action", ""),
            it.get("responsible", ""), it.get("priority", "Media"), it.get("status", "Pendiente")
        ])
    return output.getvalue().encode("utf-8-sig")

def generate_csv_action_plan(items: List[Dict[str, Any]]) -> bytes:
    import csv
    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow([
        "Codigo_Req", "Capitulo", "Requisito_con_Brecha", "Puntuacion_Actual",
        "Hallazgo_Causa_Raiz", "Accion_Correctiva_de_Mejora",
        "Responsable_Institucional", "Prioridad", "Estado_de_Ejecucion"
    ])
    for it in items:
        sc = it.get("score", 0)
        if sc in [0, 1, "0", "1"]:
            writer.writerow([
                it.get("code", ""), it.get("group_id", ""), it.get("title", ""),
                "No Cumple (0)" if sc in [0, "0"] else "Cumple Parcial (1)",
                it.get("findings", it.get("default_findings", "")),
                it.get("action", it.get("default_action", "")),
                it.get("responsible", it.get("default_responsible", "")),
                it.get("priority", it.get("default_priority", "Media")),
                it.get("status", "Pendiente")
            ])
    return output.getvalue().encode("utf-8-sig")

def generate_json_export(evaluation: Dict[str, Any], items: List[Dict[str, Any]]) -> bytes:
    data = {
        "metadata": {
            "institucion": "Optimus Health IPS",
            "sede": "Sincelejo, Sucre",
            "programa": "NEWCARE (Insuficiencias Organicas Cronicas Complejas)",
            "ente_certificador": "Fundacion New Health",
            "fecha_diagnostico": datetime.date.today().isoformat(),
            "version_herramienta": "2.0"
        },
        "resumen_diagnostico": {
            "cumplimiento_global_pct": evaluation["global_compliance"],
            "brecha_global_pct": evaluation["global_gap"],
            "meta_certificacion_pct": evaluation["target_pct"],
            "veredicto": evaluation["verdict_title"],
            "total_requisitos": evaluation["total_items"],
            "total_completo": evaluation["total_completo"],
            "total_parcial": evaluation["total_parcial"],
            "total_no_cumple": evaluation["total_no_cumple"],
            "total_na": evaluation["total_na"]
        },
        "capitulos": evaluation["groups"],
        "requisitos_evaluados": items
    }
    return json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")