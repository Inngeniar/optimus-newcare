import openpyxl

wb = openpyxl.load_workbook('Baremacion_Personal_Asistencial_NEWCARE_Optimus_Health.xlsx')

print('Workbook verification:')
for sname in wb.sheetnames:
    ws = wb[sname]
    print(f' - Sheet "{sname}": {ws.max_row} rows, {ws.max_column} cols, CF rules: {len(ws.conditional_formatting)}')

ws6 = wb['6. Registro y Evaluador']
print('\nHoja 6 Row 5 to 7 formula checks:')
for r in range(5, 8):
    cand = ws6.cell(r, 2).value
    f_d = ws6.cell(r, 4).value
    f_e = ws6.cell(r, 5).value
    f_f = ws6.cell(r, 6).value
    f_g = ws6.cell(r, 7).value
    f_h = ws6.cell(r, 8).value
    f_i = ws6.cell(r, 9).value
    print(f'R{r} ({cand}):')
    print(f'   D: {f_d}')
    print(f'   E: {f_e}')
    print(f'   F: {f_f}')
    print(f'   G: {f_g}')
    print(f'   H: {f_h}')
    print(f'   I: {f_i}')
